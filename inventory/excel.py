from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from re import fullmatch
from zipfile import ZipFile

from defusedxml.ElementTree import iterparse
from django.core.exceptions import ValidationError
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string

from .models import Document, Line, Product
from .services import validate_line_details

HEADERS = ["ma_hang", "so_luong", "don_gia"]
DETAIL_HEADERS = ["sl_chung_tu", "sl_khong_nhan", "ly_do"]
WAREHOUSE_HEADERS = ["ma_kho"]
MAX_RECEIPT_ROWS = 1000
LEGACY_RECEIPT_ROWS = 50
RECEIPT_LIMIT_ERROR = (
    "Mẫu mới tối đa 1.000 dòng; mẫu cũ tối đa 50 mã. "
    "Mỗi phiếu sẽ được chia tối đa 50 mã."
)
RECEIPT_READ_ERROR = (
    "Không đọc được Excel. File có thể bị hỏng; hãy chép dữ liệu vào mẫu .xlsx mới."
)


def _open_receipt(file):
    """Bound the real worksheet, including files with incorrect dimensions."""
    workbook = None
    try:
        file.seek(0)
        raw = file.read(5 * 1024 * 1024 + 1)
        if len(raw) > 5 * 1024 * 1024:
            raise ValidationError("File tối đa 5 MB.")
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or len({entry.filename for entry in entries}) != len(entries)
                or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024
                or any(entry.flag_bits & 1 for entry in entries)
                or any("vbaproject" in entry.filename.lower() for entry in entries)
            ):
                raise ValidationError(
                    "Cấu trúc Excel quá lớn hoặc không được hỗ trợ. Hãy dùng mẫu .xlsx mới."
                )
            workbook = load_workbook(
                BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            sheet = workbook.worksheets[0]
            with archive.open(sheet._worksheet_path) as stream:
                for _event, element in iterparse(stream, events=("start",)):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "row":
                        ref = element.attrib.get("r", "")
                        if ref and (
                            not ref.isascii()
                            or not ref.isdigit()
                            or len(ref) > 4
                            or not 1 <= int(ref) <= MAX_RECEIPT_ROWS + 1
                        ):
                            raise ValidationError(RECEIPT_LIMIT_ERROR)
                    elif tag == "c":
                        ref = element.attrib.get("r", "")
                        if ref:
                            match = fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", ref)
                            if not match:
                                raise ValidationError(RECEIPT_READ_ERROR)
                            if int(match[2]) > MAX_RECEIPT_ROWS + 1:
                                raise ValidationError(RECEIPT_LIMIT_ERROR)
                            if column_index_from_string(match[1]) > 7:
                                raise ValidationError(
                                    "File có quá nhiều cột. Chỉ dùng các cột của mẫu."
                                )
                    element.clear()
        return workbook
    except Exception as exc:
        if workbook is not None:
            workbook.close()
        if isinstance(exc, ValidationError):
            raise
        raise ValidationError(RECEIPT_READ_ERROR) from exc


def parse_receipt(file):
    errors = []
    workbook = _open_receipt(file)
    try:
        sheet = workbook.worksheets[0]
        if sheet.max_row and sheet.max_row > MAX_RECEIPT_ROWS + 1:
            raise ValidationError(RECEIPT_LIMIT_ERROR)
        if sheet.max_column and sheet.max_column > 7:
            raise ValidationError("File có quá nhiều cột. Chỉ dùng các cột của mẫu.")
        # Include one overflow row, even if a workbook understates its dimensions.
        try:
            rows = list(sheet.iter_rows(max_row=MAX_RECEIPT_ROWS + 2, max_col=8))
        except Exception as exc:
            raise ValidationError(RECEIPT_READ_ERROR) from exc
        header = rows.pop(0) if rows else None
        values = [str(c.value or "").strip() for c in header] if header else []
        multi = values[:7] == HEADERS + DETAIL_HEADERS + WAREHOUSE_HEADERS and not any(
            values[7:]
        )
        extended = values[:6] == HEADERS + DETAIL_HEADERS and not any(
            values[7:] if multi else values[6:]
        )
        legacy = values[:3] == HEADERS and not any(values[3:])
        if not (multi or extended or legacy):
            raise ValidationError(
                "Dùng mẫu nhập kho mới gồm ma_hang, so_luong, don_gia, sl_chung_tu, sl_khong_nhan, ly_do, ma_kho; vẫn nhận mẫu 3 hoặc 6 cột cũ."
            )
        products = {
            product.sku: product
            for product in Product.objects.filter(
                sku__in={str(cells[0].value or "").strip() for cells in rows}
            )
        }
        result = []
        seen = set()
        used = 0
        for number, cells in enumerate(rows, start=2):
            if all(c.value is None for c in cells):
                continue
            used += 1
            if (
                used > (MAX_RECEIPT_ROWS if multi else LEGACY_RECEIPT_ROWS)
                or number > MAX_RECEIPT_ROWS + 1
            ):
                raise ValidationError(RECEIPT_LIMIT_ERROR)
            if any(c.data_type in {"f", "e"} for c in cells) or any(
                c.value is not None
                for c in cells[7 if multi else 6 if extended else 3 :]
            ):
                errors.append(f"Dòng {number}: không dùng công thức, hãy dán giá trị.")
                continue
            raw_sku = cells[0].value
            if not isinstance(raw_sku, str) or not raw_sku.strip():
                errors.append(
                    f"Dòng {number}: mã hàng phải là văn bản (Text), không để trống; giữ nguyên số 0 đầu mã."
                )
                continue
            sku = raw_sku.strip()
            warehouse_code = str(cells[6].value or "").strip() if multi else ""
            if (
                multi
                and cells[6].value is not None
                and not isinstance(cells[6].value, str)
            ):
                errors.append(f"Dòng {number}: mã kho phải là văn bản (Text).")
                continue
            identity = (warehouse_code, sku)
            if identity in seen:
                errors.append(f"Dòng {number}: mã {sku} bị trùng tại cùng kho.")
                continue
            seen.add(identity)
            product = products.get(sku)
            if not product:
                errors.append(
                    f"Dòng {number}: mã hàng {sku or '(trống)'} chưa có trong danh mục."
                )
                continue
            try:
                if isinstance(cells[1].value, bool) or isinstance(cells[2].value, bool):
                    raise ValueError
                q = Decimal(str(cells[1].value))
                price = Decimal(str(cells[2].value))
                if (
                    not q.is_finite()
                    or not price.is_finite()
                    or q < 0
                    or price < 0
                    or q >= Decimal("1e13")
                    or price >= Decimal("1e14")
                ):
                    raise ValueError
                if q != q.quantize(Decimal(".001")) or price != price.quantize(
                    Decimal(".0001")
                ):
                    raise ValueError
                if q * price >= Decimal("1e20"):
                    raise ValueError
            except (InvalidOperation, ValueError, TypeError):
                errors.append(
                    f"Dòng {number}: số lượng phải dương (tối đa 3 số lẻ), đơn giá không âm (tối đa 4 số lẻ). Nhập ô số, ví dụ 600000; không nhập chữ 'đ', 'VND' hoặc dấu phân cách hàng nghìn trong ô văn bản."
                )
                continue
            if not extended and any(c.value is not None for c in cells[3:6]):
                errors.append(f"Dòng {number}: có dữ liệu ngoài cột mẫu.")
                continue
            try:
                quantities = []
                for cell, default in [(cells[3], None), (cells[4], Decimal(0))]:
                    value = (
                        Decimal(str(cell.value)) if cell.value is not None else default
                    )
                    if value is not None and (
                        isinstance(cell.value, bool)
                        or not value.is_finite()
                        or value < 0
                        or value >= Decimal("1e13")
                        or value != value.quantize(Decimal(".001"))
                    ):
                        raise ValueError
                    quantities.append(value)
                expected, rejected = quantities
                note = str(cells[5].value or "").strip()
                if len(note) > 250:
                    raise ValueError
                validate_line_details(
                    Document(kind="IN"),
                    Line(
                        product=product,
                        quantity=q,
                        expected_quantity=expected,
                        rejected_quantity=rejected,
                        note=note,
                    ),
                )
                if q == 0 and not (expected or rejected):
                    raise ValueError
            except (InvalidOperation, ValueError, ValidationError) as exc:
                reason = (
                    " ".join(exc.messages)
                    if isinstance(exc, ValidationError)
                    else "Số lượng chi tiết không hợp lệ hoặc lý do dài quá 250 ký tự."
                )
                errors.append(f"Dòng {number}: {reason}")
                continue
            result.append(
                {
                    "product_id": product.pk,
                    "unit": product.unit,
                    "warehouse_code": warehouse_code,
                    "row_number": number,
                    "sku": product.sku,
                    "name": product.name,
                    "quantity": str(q),
                    "unit_price": str(price),
                    "amount": str(
                        (q * price).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)
                    ),
                    "expected_quantity": str(expected)
                    if expected is not None
                    else None,
                    "rejected_quantity": str(rejected),
                    "shortage": str(max(expected - q - rejected, 0))
                    if expected is not None
                    else None,
                    "note": note,
                }
            )
        if errors:
            raise ValidationError(
                errors[:30]
                + (
                    [
                        f"File có {len(errors)} lỗi, đang hiển thị 30 lỗi đầu. Chưa tạo phiếu nào."
                    ]
                    if len(errors) > 30
                    else []
                )
            )
        if not result:
            raise ValidationError("File chưa có dòng hàng hợp lệ.")
        return result
    finally:
        workbook.close()


def receipt_template(*, user=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Nhap kho"
    sheet.append(HEADERS + DETAIL_HEADERS + WAREHOUSE_HEADERS)
    for column in ["D", "E", "F"]:
        sheet.column_dimensions[column].width = 28
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 25
    sheet.column_dimensions["G"].width = 25
    for number in range(2, MAX_RECEIPT_ROWS + 2):
        sheet.cell(number, 1).number_format = "@"
        sheet.cell(number, 7).number_format = "@"
        for column in (2, 4, 5):
            sheet.cell(number, column).number_format = "0.###"
        sheet.cell(number, 3).number_format = "#,##0.####"
    sheet.freeze_panes = "A2"
    guide = book.create_sheet("Huong dan")
    guide.column_dimensions["A"].width = 110
    for text in [
        "Điền ma_kho để phân bổ từng dòng về đúng kho. Tra mã trong trang Danh sach kho.",
        "Mã hàng phải có trong Danh mục hàng hóa. Cùng mã có thể nhập ở các kho khác nhau; không lặp cùng mã trong cùng kho.",
        "Tối đa 1.000 dòng dữ liệu ở mẫu mới (không tính dòng tiêu đề). Mỗi kho tạo phiếu riêng, trên 50 mã tự chia phiếu. Duyệt xong mới cộng tồn.",
        "Nếu bỏ trống ma_kho, dòng đó dùng Kho nhận chung đã chọn trên màn hình. Nếu không chọn kho chung, mọi dòng phải có ma_kho.",
        "so_luong là số thực nhận đạt yêu cầu; phần thiếu/hỏng ghi sl_chung_tu, sl_khong_nhan và ly_do. Đơn giá miễn phí ghi 0.",
        "Số lượng và đơn giá dùng ô số: nhập 600000, có thể định dạng hiển thị tiền; không dán chuỗi chữ như 600.000 đ. Không để dòng chỉ có đơn giá 0 mà thiếu mã hàng/số lượng.",
        "Không dùng công thức. Ngày nghiệp vụ là hôm nay. Chỉ dùng các kho bạn được phân quyền.",
    ]:
        guide.append([text])
    warehouses = book.create_sheet("Danh sach kho")
    warehouses.append(["ma_kho", "ten_kho"])
    warehouses.column_dimensions["A"].width = 25
    warehouses.column_dimensions["B"].width = 50
    if user is not None:
        from .services import warehouses_for

        for warehouse in warehouses_for(user).order_by("code"):
            warehouses.append([warehouse.code, warehouse.name])
            warehouses.cell(warehouses.max_row, 1).number_format = "@"
            for cell in warehouses[warehouses.max_row]:
                cell.data_type = "s"
    return book
