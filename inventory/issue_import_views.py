"""Excel is an alternate input for an OUT draft, never an approval shortcut."""

import time
from datetime import date
from decimal import Decimal, InvalidOperation
from uuid import uuid4
from zipfile import ZipFile

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_GET, require_http_methods
from openpyxl import Workbook, load_workbook

from .access import require_function
from .forms import IssueImportForm
from .models import Balance, Document, Line, Product, Warehouse
from .services import audit_document, capture_lines, validate_document, warehouses_for

HEADERS = ["ma_hang", "so_luong", "ghi_chu"]
PREVIEW_SECONDS = 1800


def parse_issue(upload):
    try:
        with ZipFile(upload) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024
            ):
                raise ValidationError(
                    "Nội dung giải nén quá lớn. Hãy dùng file mẫu mới."
                )
        upload.seek(0)
        book = load_workbook(upload, read_only=True, data_only=False, keep_links=False)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(
            "Không đọc được Excel. Hãy lưu file .xlsx theo mẫu."
        ) from exc
    try:
        sheet = book.worksheets[0]
        if (sheet.max_row and sheet.max_row > 1000) or (
            sheet.max_column and sheet.max_column > 3
        ):
            raise ValidationError("Dùng file mẫu mới: tối đa 50 mã và đúng 3 cột.")
        iterator = sheet.iter_rows(max_row=1001, max_col=4)
        header = next(iterator)
        if [str(c.value or "").strip() for c in header[:3]] != HEADERS or header[
            3
        ].value is not None:
            raise ValidationError("Dùng đúng mẫu xuất kho: ma_hang, so_luong, ghi_chu.")
        rows, errors, seen = [], [], set()
        used_rows = 0
        for number, cells in enumerate(iterator, 2):
            if all(c.value is None for c in cells):
                continue
            used_rows += 1
            if used_rows > 50 or number > 1000:
                raise ValidationError("Mỗi phiếu tối đa 50 mã hàng. Hãy chia file.")
            if (
                any(c.data_type in {"f", "e"} for c in cells)
                or cells[3].value is not None
            ):
                errors.append(
                    f"Dòng {number}: không dùng công thức, lỗi Excel hoặc cột ngoài mẫu."
                )
                continue
            sku = cells[0].value
            if not isinstance(sku, str) or not sku.strip() or len(sku.strip()) > 100:
                errors.append(
                    f"Dòng {number}: mã hàng phải là văn bản; giữ nguyên số 0 đầu mã."
                )
                continue
            sku = sku.strip()
            if sku in seen:
                errors.append(f"Dòng {number}: mã {sku} bị trùng.")
                continue
            seen.add(sku)
            try:
                quantity = Decimal(str(cells[1].value))
                if (
                    isinstance(cells[1].value, bool)
                    or not quantity.is_finite()
                    or not (0 < quantity < Decimal("1e13"))
                    or quantity != quantity.quantize(Decimal(".001"))
                ):
                    raise ValueError
            except (InvalidOperation, ValueError, TypeError):
                errors.append(f"Dòng {number}: số lượng phải dương, tối đa 3 số lẻ.")
                continue
            note = str(cells[2].value or "").strip()
            if len(note) > 250:
                errors.append(f"Dòng {number}: ghi chú tối đa 250 ký tự.")
                continue
            rows.append({"sku": sku, "quantity": str(quantity), "note": note})
        products = {p.sku: p for p in Product.objects.filter(sku__in=seen)}
        for row in rows:
            product = products.get(row["sku"])
            if product is None:
                errors.append(f"Mã {row['sku']}: chưa có trong Hàng hóa.")
            else:
                row.update(product_id=product.pk, name=product.name, unit=product.unit)
        if errors:
            raise ValidationError(errors[:30])
        if not rows:
            raise ValidationError("File chưa có dòng hàng hợp lệ.")
        return rows
    finally:
        book.close()


def check_stock(warehouse, rows):
    balances = {
        b.product_id: b.quantity
        for b in Balance.objects.filter(
            warehouse=warehouse, product_id__in=[r["product_id"] for r in rows]
        )
    }
    products = dict(
        Product.objects.filter(pk__in=[r["product_id"] for r in rows]).values_list(
            "pk", "sku"
        )
    )
    errors = []
    for row in rows:
        if products.get(row["product_id"]) != row["sku"]:
            errors.append(f"Mã {row['sku']}: danh mục đã thay đổi. Hãy tải lại file.")
            continue
        available = balances.get(row["product_id"], Decimal(0))
        row["available"] = str(available)
        if Decimal(row["quantity"]) > available:
            errors.append(
                f"Mã {row['sku']}: yêu cầu {row['quantity']}, tồn có thể xuất chỉ {available}."
            )
    if errors:
        raise ValidationError(errors)


@login_required
@require_GET
def issue_template(request):
    require_function(request.user, "issues_create")
    book = Workbook()
    sheet = book.active
    sheet.title = "Xuat kho"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    for column, width in [("A", 25), ("B", 20), ("C", 45)]:
        sheet.column_dimensions[column].width = width
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="mau-xuat-kho.xlsx"'
    book.save(response)
    return response


@login_required
@require_http_methods(["GET", "POST"])
def issue_import(request):
    require_function(request.user, "issues_create")
    form = IssueImportForm(
        request.POST or None, request.FILES or None, user=request.user
    )
    preview = None
    if request.method == "POST" and request.POST.get("confirm"):
        candidate = request.session.get("issue_import")
        if (
            not candidate
            or candidate.get("owner") != request.user.pk
            or request.POST.get("import_key") != candidate.get("key")
            or not 0 <= time.time() - candidate.get("created_at", 0) <= PREVIEW_SECONDS
        ):
            request.session.pop("issue_import", None)
            messages.error(
                request, "Bản xem trước đã hết hạn hoặc đã tạo phiếu. Hãy tải lại file."
            )
            return redirect("issue_import")
        preview = candidate
        try:
            with transaction.atomic():
                actor = (
                    get_user_model().objects.select_for_update().get(pk=request.user.pk)
                )
                require_function(actor, "issues_create")
                warehouse = get_object_or_404(
                    warehouses_for(actor), pk=preview["warehouse_id"]
                )
                Warehouse.objects.select_for_update().get(pk=warehouse.pk)
                check_stock(warehouse, preview["rows"])
                details = {
                    key: date.fromisoformat(value)
                    if key in {"business_date", "reference_date", "due_date"} and value
                    else value
                    for key, value in preview["details"].items()
                }
                doc = Document(
                    kind="OUT",
                    source=warehouse,
                    created_by=actor,
                    note=preview["note"],
                    import_key=preview["key"],
                    **details,
                )
                validate_document(doc)
                existing = Document.objects.filter(import_key=preview["key"]).first()
                if existing:
                    doc = existing
                else:
                    doc.save()
                    Line.objects.bulk_create(
                        capture_lines(
                            doc,
                            [
                                Line(
                                    document=doc,
                                    product_id=row["product_id"],
                                    quantity=Decimal(row["quantity"]),
                                    note=row["note"],
                                )
                                for row in preview["rows"]
                            ],
                        )
                    )
                    audit_document(doc, actor, "CREATE")
            request.session.pop("issue_import", None)
            messages.success(
                request,
                "Đã tạo phiếu xuất nháp. Chưa trừ tồn; người có quyền cần kiểm tra và duyệt.",
            )
            return redirect("document_detail", pk=doc.pk)
        except ValidationError as exc:
            messages.error(request, " ".join(exc.messages))
            request.session.pop("issue_import", None)
            return redirect("issue_import")
    elif request.method == "POST":
        request.session.pop("issue_import", None)
        if form.is_valid():
            try:
                rows = parse_issue(form.cleaned_data["file"])
                warehouse = form.cleaned_data["warehouse"]
                check_stock(warehouse, rows)
                preview = {
                    "key": str(uuid4()),
                    "owner": request.user.pk,
                    "created_at": time.time(),
                    "warehouse_id": warehouse.pk,
                    "warehouse_name": str(warehouse),
                    "note": form.cleaned_data["note"],
                    "rows": rows,
                    "details": {
                        key: str(form.cleaned_data[key])
                        if isinstance(form.cleaned_data[key], date)
                        else form.cleaned_data[key]
                        for key in form.Meta.fields
                        if key != "note"
                    },
                }
                request.session["issue_import"] = preview
            except ValidationError as exc:
                form.add_error("file", exc)
    return render(request, "issue_import.html", {"form": form, "preview": preview})
