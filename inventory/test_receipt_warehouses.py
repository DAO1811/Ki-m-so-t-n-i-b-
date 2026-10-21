from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import timedelta
from io import BytesIO
from re import sub
from threading import Barrier
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections, connection, connections
from django.test import Client, TestCase, TransactionTestCase, skipUnlessDBFeature
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .excel import DETAIL_HEADERS, HEADERS, WAREHOUSE_HEADERS
from .models import Balance, Document, Line, Movement, Product, Warehouse
from .receipt_import_views import confirm_receipts
from .services import post_document
from .test_product_import import workbook_upload
from .test_support import grant_functions

COLUMNS = HEADERS + DETAIL_HEADERS + WAREHOUSE_HEADERS


class ReceiptWarehouseTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.maker = User.objects.create_user("receipt-file-maker")
        cls.head = User.objects.create_user("receipt-file-head")
        grant_functions(cls.maker, {"receipts_create"}, level=10)
        grant_functions(cls.head, {"receipts_approve"}, level=30)
        cls.a = Warehouse.objects.create(code="K01", name="Kho A")
        cls.b = Warehouse.objects.create(code="K02", name="Kho B")
        cls.outside = Warehouse.objects.create(code="HIDDEN", name="Không được cấp")
        for warehouse in [cls.a, cls.b]:
            warehouse.members.add(cls.maker)
            warehouse.managers.add(cls.head)
        cls.p = Product.objects.create(sku="CAM001", name="Camera", unit="Cái")
        cls.url = reverse("receipt_import")

    def setUp(self):
        self.client.force_login(self.maker)

    def upload(self, rows=None, *, warehouse=None, headers=COLUMNS, **kwargs):
        data = {
            "file": workbook_upload(
                rows or [["CAM001", 2, 10, None, None, None, "K01"]], headers
            ),
            "note": "Nhập vật tư",
        }
        if warehouse is not None:
            data["warehouse"] = warehouse.pk
        data.update(kwargs)
        return self.client.post(self.url, data)

    def confirm(self):
        key = self.client.session["receipt_import"]["key"]
        return self.client.post(self.url, {"confirm": 1, "import_key": key})

    def test_actual_receipt_download_has_warehouse_column_and_authorized_codes(self):
        response = self.client.get(reverse("receipt_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("mau-nhap-kho-theo-kho.xlsx", response["Content-Disposition"])
        self.assertIn("no-store", response["Cache-Control"])
        book = load_workbook(BytesIO(response.content))
        self.assertEqual([c.value for c in book.worksheets[0][1]], COLUMNS)
        self.assertEqual(
            set(row[0] for row in list(book["Danh sach kho"].values)[1:]),
            {"K01", "K02"},
        )
        self.assertEqual(book.worksheets[0]["G2"].number_format, "@")
        self.assertEqual(book.worksheets[0]["A1001"].number_format, "@")
        self.assertEqual(book.worksheets[0]["G1001"].number_format, "@")
        self.assertEqual(book.worksheets[0].max_row, 1001)
        self.assertIn(
            "1.000 dòng", " ".join(str(r[0]) for r in book["Huong dan"].values)
        )
        book.close()

    def test_existing_product_to_multiple_warehouses_without_common_warehouse(self):
        response = self.upload(
            [
                ["CAM001", 2, 10, None, None, None, "K01"],
                ["CAM001", 3, 20, None, None, None, "K02"],
            ]
        )
        self.assertContains(response, "Kho A")
        self.assertContains(response, "Kho B")
        self.assertEqual(self.client.session["receipt_import"]["document_count"], 2)
        self.assertFalse(Document.objects.exists())
        self.confirm()
        self.assertEqual(Product.objects.count(), 1)
        self.assertEqual(Document.objects.count(), 2)
        self.assertFalse(Balance.objects.exists())
        self.assertFalse(Movement.objects.exists())
        for document in Document.objects.all():
            self.assertEqual(document.status, "DRAFT")
            post_document(document.pk, self.head)
        self.assertEqual(
            (
                Balance.objects.get(warehouse=self.a).quantity,
                Balance.objects.get(warehouse=self.b).value,
            ),
            (2, 60),
        )

    def test_file_warehouse_wins_and_blank_rows_use_common_warehouse(self):
        self.upload(
            [
                ["CAM001", 2, 10, None, None, None, None],
                ["CAM001", 3, 20, None, None, None, "K02"],
            ],
            warehouse=self.a,
        )
        self.confirm()
        self.assertEqual(
            set(Document.objects.values_list("destination_id", flat=True)),
            {self.a.pk, self.b.pk},
        )

    def test_blank_and_explicit_same_warehouse_duplicate_is_blocked(self):
        self.upload(
            [
                ["CAM001", 1, 10, None, None, None, None],
                ["CAM001", 2, 10, None, None, None, "K01"],
            ],
            warehouse=self.a,
        )
        self.assertNotIn("receipt_import", self.client.session)
        self.assertFalse(Document.objects.exists())

    def test_no_fallback_unknown_outside_numeric_warehouse_and_formula_block_entire_file(
        self,
    ):
        for code in [None, "MISSING", "HIDDEN", 1, "=1+1"]:
            with self.subTest(code=code):
                self.upload(
                    [
                        ["CAM001", 1, 10, None, None, None, "K01"],
                        ["CAM001", 1, 10, None, None, None, code],
                    ]
                )
                self.assertNotIn("receipt_import", self.client.session)
                self.assertFalse(Document.objects.exists())

    def test_partial_damage_preserved_and_only_accepted_quantity_is_posted(self):
        self.upload([["CAM001", 7, 10, 10, 2, "Thiếu 1, hỏng 2", "K01"]])
        self.confirm()
        line = Line.objects.get()
        self.assertEqual(
            (
                line.quantity,
                line.expected_quantity,
                line.rejected_quantity,
                line.shortage,
            ),
            (7, 10, 2, 1),
        )
        post_document(Document.objects.get().pk, self.head)
        self.assertEqual(Balance.objects.get().quantity, 7)

    def test_revoked_warehouse_and_changed_product_reject_without_partial_receipts(
        self,
    ):
        self.upload()
        self.a.members.remove(self.maker)
        self.confirm()
        self.assertFalse(Document.objects.exists())
        self.a.members.add(self.maker)
        self.upload()
        Product.objects.filter(pk=self.p.pk).update(unit="Bộ")
        self.confirm()
        self.assertFalse(Document.objects.exists())

    def test_warehouse_code_renamed_after_preview_is_not_silently_used(self):
        self.upload()
        Warehouse.objects.filter(pk=self.a.pk).update(code="RENAMED")
        self.confirm()
        self.assertFalse(Document.objects.exists())

    def test_expired_preview_and_day_rollover_require_upload_again(self):
        for field in ["created_at", "business_date"]:
            self.upload()
            session = self.client.session
            preview = session["receipt_import"]
            if field == "created_at":
                preview[field] = 0
            else:
                preview["details"][field] = str(
                    timezone.localdate() - timedelta(days=1)
                )
            session["receipt_import"] = preview
            session.save()
            self.confirm()
            self.assertFalse(Document.objects.exists())

    def test_second_receipt_failure_rolls_back_first_and_lines(self):
        self.upload(
            [
                ["CAM001", 2, 10, None, None, None, "K01"],
                ["CAM001", 3, 20, None, None, None, "K02"],
            ]
        )
        from .services import audit_document as real_audit

        calls = []

        def audit(document, actor, action):
            calls.append(document.pk)
            if len(calls) == 2:
                raise ValidationError("second receipt failed")
            return real_audit(document, actor, action)

        with patch("inventory.receipt_import_views.audit_document", side_effect=audit):
            self.confirm()
        self.assertEqual(len(calls), 2)
        self.assertFalse(Document.objects.exists())
        self.assertFalse(Line.objects.exists())

    def test_replay_and_direct_repeated_confirmation_do_not_duplicate_receipts(self):
        self.upload(
            [
                ["CAM001", 2, 10, None, None, None, "K01"],
                ["CAM001", 3, 20, None, None, None, "K02"],
            ]
        )
        preview = deepcopy(self.client.session["receipt_import"])
        result = confirm_receipts(self.maker, preview)
        again = confirm_receipts(self.maker, preview)
        self.assertEqual([d.pk for d in result], [d.pk for d in again])
        self.confirm()
        self.assertEqual(Document.objects.count(), 2)
        self.client.post(self.url, {"confirm": 1, "import_key": preview["key"]})
        self.assertEqual(Document.objects.count(), 2)
        key = self.client.session["receipt_import_result"]["key"]
        response = self.client.get(
            reverse("documents") + "?kind=IN&import_batch=" + key
        )
        self.assertEqual(response.context["page_obj"].paginator.count, 2)

    def large_rows(self, count):
        Product.objects.bulk_create(
            [Product(sku=f"R{i}", name=f"Hàng {i}") for i in range(count)]
        )
        return [
            [f"R{i}", 1, 10, None, None, None, "K01" if i < 550 else "K02"]
            for i in range(count)
        ]

    def test_new_template_1000_rows_split_into_20_drafts_by_warehouse(self):
        self.assertContains(self.client.get(self.url), "5 MB / 1.000 dòng")
        response = self.upload(self.large_rows(1000))
        self.assertContains(response, "1000 dòng hợp lệ")
        preview = self.client.session["receipt_import"]
        self.assertEqual(len(preview["rows"]), 1000)
        self.assertEqual(preview["rows"][-1]["sku"], "R999")
        self.assertEqual(preview["rows"][-1]["row_number"], 1001)
        self.assertEqual(preview["document_count"], 20)
        self.confirm()
        self.assertEqual(Document.objects.filter(destination=self.a).count(), 11)
        self.assertEqual(Document.objects.filter(destination=self.b).count(), 9)
        self.assertEqual(Line.objects.count(), 1000)
        self.assertTrue(all(d.lines.count() == 50 for d in Document.objects.all()))
        self.assertFalse(Document.objects.exclude(status="DRAFT").exists())
        self.assertFalse(Balance.objects.exists())
        self.assertFalse(Movement.objects.exists())

    def test_1001_rows_rejected_even_with_understated_excel_dimensions(self):
        rows = self.large_rows(1001)
        for understated in (False, True):
            with self.subTest(understated=understated):
                upload = workbook_upload(rows, COLUMNS)
                if understated:
                    output = BytesIO()
                    with (
                        ZipFile(upload) as source,
                        ZipFile(output, "w", ZIP_DEFLATED) as target,
                    ):
                        for entry in source.infolist():
                            data = source.read(entry.filename)
                            if entry.filename == "xl/worksheets/sheet1.xml":
                                data = sub(
                                    rb'<dimension ref="[^"]+"',
                                    b'<dimension ref="A1:G2"',
                                    data,
                                )
                            target.writestr(entry, data)
                    upload = SimpleUploadedFile("overflow.xlsx", output.getvalue())
                response = self.upload(file=upload)
                self.assertContains(response, "Mẫu mới tối đa 1.000 dòng")
                self.assertNotIn("receipt_import", self.client.session)
                self.assertFalse(Document.objects.exists())

    def test_last_of_1000_rows_is_validated_without_partial_preview(self):
        rows = self.large_rows(1000)
        rows[-1][0] = "MISSING-LAST-PRODUCT"
        response = self.upload(rows)
        self.assertContains(response, "Dòng 1001")
        self.assertContains(response, "MISSING-LAST-PRODUCT")
        self.assertNotIn("receipt_import", self.client.session)
        self.assertFalse(Document.objects.exists())

    def test_legacy_three_and_six_column_files_still_limit_to_50_rows(self):
        rows = self.large_rows(51)
        for headers in (HEADERS, HEADERS + DETAIL_HEADERS):
            with self.subTest(headers=headers):
                response = self.upload(
                    [r[: len(headers)] for r in rows], warehouse=self.a, headers=headers
                )
                self.assertContains(response, "mẫu cũ tối đa 50 mã")
                self.assertNotIn("receipt_import", self.client.session)
                self.assertFalse(Document.objects.exists())

    def rewritten_upload(self, upload, rewrite):
        output = BytesIO()
        with ZipFile(upload) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    data = rewrite(data)
                target.writestr(entry, data)
        return SimpleUploadedFile("receipt.xlsx", output.getvalue())

    def test_numeric_sku_is_not_silently_converted_and_text_preserves_leading_zero(
        self,
    ):
        Product.objects.create(sku="123", name="Numeric-looking code")
        Product.objects.create(sku="00123", name="Leading-zero code")
        response = self.upload([[123, 1, 10, None, None, None, "K01"]])
        self.assertContains(response, "mã hàng phải là văn bản (Text)")
        self.assertNotIn("receipt_import", self.client.session)
        self.upload([["00123", 1, 10, None, None, None, "K01"]])
        self.confirm()
        self.assertEqual(Line.objects.get().sku_snapshot, "00123")

    def test_currency_text_and_orphan_zero_row_do_not_produce_partial_preview(self):
        for bad in (
            ["CAM001", 1, "600.000 đ", None, None, None, "K02"],
            [None, None, 0, None, None, None, None],
        ):
            with self.subTest(bad=bad):
                response = self.upload(
                    [["CAM001", 1, 10, None, None, None, "K01"], bad]
                )
                self.assertContains(response, "Dòng 3")
                self.assertNotIn("receipt_import", self.client.session)
                self.assertFalse(Document.objects.exists())
                self.assertFalse(Balance.objects.exists())

    def test_many_errors_report_total_instead_of_hiding_remaining_invalid_rows(self):
        rows = self.large_rows(31)
        for row in rows:
            row[2] = "600.000 đ"
        response = self.upload(rows)
        self.assertContains(response, "File có 31 lỗi, đang hiển thị 30 lỗi đầu")
        self.assertContains(response, "ví dụ 600000")
        self.assertNotIn("receipt_import", self.client.session)
        self.assertFalse(Document.objects.exists())

    def test_truncated_xml_and_invalid_numeric_cell_show_file_error_instead_of_500(
        self,
    ):
        for rewrite in (
            lambda data: data[:-20],
            lambda data: sub(
                rb'(<c r="B2"[^>]*><v>)[^<]*(</v>)',
                rb"\g<1>not-a-number\2",
                data,
            ),
        ):
            with self.subTest(rewrite=rewrite):
                upload = self.rewritten_upload(
                    workbook_upload(
                        [["CAM001", 1, 10, None, None, None, "K01"]], COLUMNS
                    ),
                    rewrite,
                )
                response = self.upload(file=upload)
                self.assertContains(response, "Không đọc được Excel")
                self.assertNotIn("receipt_import", self.client.session)
                self.assertFalse(Document.objects.exists())

    def test_hidden_extra_column_or_far_row_cannot_be_silently_discarded(self):
        for column, row in ((9, 2), (1, 5000)):
            with self.subTest(column=column, row=row):
                upload = workbook_upload(
                    [["CAM001", 1, 10, None, None, None, "K01"]], COLUMNS
                )
                book = load_workbook(upload)
                book.worksheets[0].cell(row, column, "unexpected data")
                stream = BytesIO()
                book.save(stream)
                book.close()
                stream.seek(0)
                upload = self.rewritten_upload(
                    stream,
                    lambda data: sub(
                        rb'<dimension ref="[^"]+"', b'<dimension ref="A1:G2"', data
                    ),
                )
                response = self.upload(file=upload)
                self.assertContains(
                    response,
                    "quá nhiều cột" if column == 9 else "Mẫu mới tối đa 1.000 dòng",
                )
                self.assertNotIn("receipt_import", self.client.session)
                self.assertFalse(Document.objects.exists())

    def test_wrong_extension_and_file_over_five_mb_do_not_keep_old_preview(self):
        for upload in (
            SimpleUploadedFile("receipt.xls", b"not xlsx"),
            SimpleUploadedFile("receipt.xlsx", b"0" * (5 * 1024 * 1024 + 1)),
        ):
            with self.subTest(name=upload.name):
                self.upload()
                response = self.upload(file=upload)
                self.assertEqual(response.status_code, 200)
                self.assertNotIn("receipt_import", self.client.session)
                self.assertFalse(Document.objects.exists())

    def test_warehouse_labels_in_downloaded_template_are_literal_text(self):
        warehouse = Warehouse.objects.create(code="=1+1", name="=2+2")
        warehouse.members.add(self.maker)
        response = self.client.get(reverse("receipt_template"))
        book = load_workbook(BytesIO(response.content), data_only=False)
        cells = next(
            row for row in book["Danh sach kho"].iter_rows() if row[0].value == "=1+1"
        )
        self.assertTrue(all(cell.data_type == "s" for cell in cells))
        self.assertEqual(cells[1].value, "=2+2")
        book.close()

    def test_confirmation_cannot_override_preview_warehouse_quantity_or_status(self):
        self.upload()
        key = self.client.session["receipt_import"]["key"]
        response = self.client.post(
            self.url,
            {
                "confirm": 1,
                "import_key": key,
                "warehouse": self.outside.pk,
                "quantity": "99999",
                "status": "POSTED",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Document.objects.get().destination, self.a)
        self.assertEqual(Document.objects.get().status, "DRAFT")
        self.assertEqual(Line.objects.get().quantity, 2)
        self.assertFalse(Balance.objects.exists())

    def test_revoked_create_permission_blocks_confirmation_and_template(self):
        self.upload()
        grant_functions(self.maker, {"receipts_view"})
        self.assertEqual(self.confirm().status_code, 403)
        self.assertEqual(self.client.get(reverse("receipt_template")).status_code, 403)
        self.assertFalse(Document.objects.exists())

    def test_confirmation_requires_csrf_token(self):
        self.upload()
        client = Client(enforce_csrf_checks=True)
        client.cookies = self.client.cookies
        response = client.post(
            self.url,
            {"confirm": 1, "import_key": self.client.session["receipt_import"]["key"]},
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(Document.objects.exists())


class ReceiptConfirmationConcurrencyTests(TransactionTestCase):
    @skipUnlessDBFeature("has_select_for_update")
    def test_parallel_confirmations_create_each_receipt_once(self):
        maker = get_user_model().objects.create_user("parallel-receipt-maker")
        grant_functions(maker, {"receipts_create"})
        for code in ("K01", "K02"):
            Warehouse.objects.create(code=code, name=code).members.add(maker)
        Product.objects.create(sku="P", name="Parallel confirmation item")
        self.client.force_login(maker)
        self.client.post(
            reverse("receipt_import"),
            {
                "file": workbook_upload(
                    [
                        ["P", 2, 10, None, None, None, "K01"],
                        ["P", 3, 10, None, None, None, "K02"],
                    ],
                    COLUMNS,
                )
            },
        )
        preview = deepcopy(self.client.session["receipt_import"])
        barrier = Barrier(2)

        def confirm():
            close_old_connections()
            try:
                actor = get_user_model().objects.get(pk=maker.pk)
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout = '5s'")
                    cursor.execute("SET statement_timeout = '15s'")
                barrier.wait(timeout=10)
                return [doc.pk for doc in confirm_receipts(actor, deepcopy(preview))]
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: confirm(), range(2)))
        self.assertEqual(results[0], results[1])
        self.assertEqual(Document.objects.count(), 2)
        self.assertEqual(Line.objects.count(), 2)
        self.assertFalse(Balance.objects.exists())
        self.assertFalse(Movement.objects.exists())
