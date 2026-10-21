from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .excel import DETAIL_HEADERS, HEADERS, WAREHOUSE_HEADERS
from .models import Balance, Document, Line, Movement, Product, Warehouse
from .test_product_import import workbook_upload
from .test_support import grant_functions

COLUMNS = HEADERS + DETAIL_HEADERS + WAREHOUSE_HEADERS


class ReceiptDownloadTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user("receipt-download-reader")
        grant_functions(cls.user, {"receipts_create"})
        cls.warehouse = Warehouse.objects.create(
            code="K01", name="Authorized warehouse"
        )
        cls.warehouse.members.add(cls.user)
        Warehouse.objects.create(code="HIDDEN", name="Not authorized")
        for sku in ("P1", "P2", "P3"):
            Product.objects.create(sku=sku, name=sku)

    def setUp(self):
        self.client.force_login(self.user)

    def download(self, rows):
        return self.client.post(
            reverse("receipt_import"),
            {"normalize": "1", "file": workbook_upload(rows, COLUMNS)},
        )

    def test_private_download_changes_only_unambiguous_prices_and_orphan_tail(self):
        response = self.download(
            [
                ["P1", 2, "600.000 đ", None, None, "Receipt one", "K01"],
                ["P2", 1, "6.400.000 VND", None, None, "Receipt two", "K01"],
                ["P3", 3, 0, None, None, None, "K01"],
                [None, None, 0, None, None, None, None],
            ]
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn("nhap-kho-da-chuan-hoa.xlsx", response["Content-Disposition"])
        self.assertIn("private, no-store", response["Cache-Control"])
        book = load_workbook(BytesIO(response.content), data_only=False)
        sheet = book.worksheets[0]
        rows = [list(row) for row in sheet.values if any(v is not None for v in row)]
        self.assertEqual(rows[0], COLUMNS)
        self.assertEqual(
            rows[1:],
            [
                ["P1", 2, 600000, None, None, "Receipt one", "K01"],
                ["P2", 1, 6400000, None, None, "Receipt two", "K01"],
                ["P3", 3, 0, None, None, None, "K01"],
            ],
        )
        self.assertEqual(sheet["C2"].data_type, "n")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertIn(
            "2 ô giá chữ sang số", " ".join(str(r[0]) for r in book["Huong dan"].values)
        )
        self.assertEqual(len(list(book["Can doi chieu"].values)), 2)
        self.assertEqual(
            list(book["Danh sach kho"].values),
            [("ma_kho", "ten_kho"), ("K01", "Authorized warehouse")],
        )
        book.close()
        self.assertNotIn("receipt_import", self.client.session)
        self.assertFalse(Document.objects.exists())
        self.assertFalse(Balance.objects.exists())
        self.assertFalse(Movement.objects.exists())
        # Download alone never creates documents. The usual preview/confirm is still required.
        preview_response = self.client.post(
            reverse("receipt_import"),
            {"file": SimpleUploadedFile("normalized.xlsx", response.content)},
        )
        self.assertEqual(preview_response.status_code, 200)
        preview = self.client.session["receipt_import"]
        self.assertEqual(len(preview["rows"]), 3)
        self.client.post(
            reverse("receipt_import"), {"confirm": "1", "import_key": preview["key"]}
        )
        self.assertEqual(Document.objects.get().status, "DRAFT")
        self.assertEqual(Line.objects.count(), 3)
        self.assertFalse(Balance.objects.exists())

    def test_ambiguous_price_and_formulas_are_not_guessed_or_executed(self):
        for price in ("1.234", "1,234 đ", "12.34 đ", "=1+1", "9" * 5000 + " đ"):
            with self.subTest(price=price[:20]):
                response = self.download([["P1", 1, price, None, None, None, "K01"]])
                self.assertEqual(response.status_code, 200)
                self.assertNotIn("Content-Disposition", response)
                self.assertFalse(Document.objects.exists())

    def test_download_is_scoped_to_receipt_permission_and_csrf(self):
        client = Client(enforce_csrf_checks=True)
        client.cookies = self.client.cookies
        response = client.post(
            reverse("receipt_import"),
            {
                "normalize": "1",
                "file": workbook_upload(
                    [["P1", 1, "600.000 đ", None, None, None, "K01"]], COLUMNS
                ),
            },
        )
        self.assertEqual(response.status_code, 403)
        grant_functions(self.user, {"stock_view"})
        self.assertEqual(
            self.download([["P1", 1, 0, None, None, None, "K01"]]).status_code, 403
        )
        self.assertFalse(Document.objects.exists())

    def test_normalized_workbook_keeps_1000_rows_and_exact_last_value(self):
        rows = [
            [f"P{i:04}", i + 1, "1.000 đ", None, None, None, "K01"] for i in range(1000)
        ]
        response = self.download(rows)
        self.assertIn("Content-Disposition", response)
        book = load_workbook(BytesIO(response.content))
        values = list(book.worksheets[0].values)
        self.assertEqual(len(values), 1001)
        self.assertEqual(values[-1], ("P0999", 1000, 1000, None, None, None, "K01"))
        book.close()
        self.assertFalse(Document.objects.exists())
