import json
from datetime import timedelta
from io import BytesIO, StringIO
from unittest import skipUnless
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.cache import caches
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db import connection
from django.test import Client, TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from psycopg import sql

from . import storage_reports as reports
from .drive import API
from .models import (
    Balance,
    Document,
    DriveConnection,
    Line,
    Movement,
    Product,
    StoredFile,
    Warehouse,
)
from .storage_audit import cache_key, measure_storage
from .test_support import grant_functions

cache = caches["storage"]


class StorageAccessTests(TestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)
        self.admin = get_user_model().objects.create_superuser(
            "storage-owner", password="1"
        )
        self.employee = get_user_model().objects.create_user(
            "storage-reader", password="1", is_staff=True
        )
        grant_functions(
            self.employee, {"products_view", "reports_view", "tools_export"}
        )

    @patch("inventory.storage_views.measure_storage")
    @patch("inventory.storage_views.export_month")
    def test_admin_only_including_download_and_post_methods(self, export, measure):
        for name in (
            "storage_settings",
            "storage_measure",
            "storage_download",
            "storage_drive_usage",
            "storage_export",
        ):
            url = reverse(name)
            self.assertEqual(self.client.get(url).status_code, 302)
            self.client.force_login(self.employee)
            self.assertEqual(self.client.get(url).status_code, 403)
            self.assertEqual(self.client.post(url).status_code, 403)
            self.client.logout()
        export.assert_not_called()
        measure.assert_not_called()
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("storage_measure")).status_code, 405)
        self.assertEqual(self.client.get(reverse("storage_export")).status_code, 405)
        csrf = Client(enforce_csrf_checks=True)
        csrf.force_login(self.admin)
        self.assertEqual(
            csrf.post(reverse("storage_export"), {"period": "2026-01"}).status_code, 403
        )

    @patch("inventory.storage_views.drive.google_request")
    def test_get_never_exports_or_requests_google_and_no_unknown_quota_guess(
        self, google
    ):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("storage_settings"))
        self.assertContains(response, "Chưa có số đo")
        self.assertContains(response, "Báo cáo quản trị / năm")
        self.assertEqual(response.headers["Cache-Control"], "private, no-store")
        self.assertEqual(response.context["current_module"]["key"], "administration")
        google.assert_not_called()
        report = measure_storage()
        self.assertIsNone(report["remaining_bytes"])
        self.assertIsNone(report["configured_limit_bytes"])
        self.assertNotIn("password", json.dumps(report))
        response = self.client.get(reverse("storage_download"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertNotIn(self.admin.password.encode(), response.content)

    @patch(
        "inventory.storage_views.export_month",
        side_effect=ValidationError("Drive chưa kết nối"),
    )
    def test_failure_keeps_source_and_does_not_claim_success(self, _export):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("storage_export"), {"period": "2026-01"}, follow=True
        )
        self.assertContains(response, "Drive chưa kết nối")
        self.assertNotContains(response, "đã đối soát")
        self.assertIsNone(cache.get(cache_key() + ":report"))

    @patch("inventory.storage_views.drive.access_token", return_value="test-token")
    @patch("inventory.storage_views.drive.google_request")
    def test_google_quota_belongs_to_same_account_and_unknown_limit_is_not_zero(
        self, google, _token
    ):
        DriveConnection.objects.create(
            account_email="owner@example.test",
            encrypted_refresh_token="not-returned",
            connected_by=self.admin,
        )
        self.client.force_login(self.admin)
        google.return_value = {
            "user": {"emailAddress": "wrong@example.test"},
            "storageQuota": {"usage": "10", "limit": "100"},
        }
        self.client.post(reverse("storage_drive_usage"))
        self.assertIsNone(cache.get(cache_key() + ":drive"))
        google.return_value = {
            "user": {"emailAddress": "owner@example.test"},
            "storageQuota": {"usage": "10"},
        }
        self.client.post(reverse("storage_drive_usage"))
        self.assertEqual(cache.get(cache_key() + ":drive")["used"], 10)
        self.assertIsNone(cache.get(cache_key() + ":drive")["remaining"])


class StorageReportsTests(TransactionTestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)
        self.admin = get_user_model().objects.create_superuser(
            "report-owner", password="1"
        )
        self.a = Warehouse.objects.create(code="A", name="Kho A")
        self.b = Warehouse.objects.create(code="B", name="Kho B")
        self.p = Product.objects.create(
            sku="00123", name='=HYPERLINK("https://invalid.test")', unit="Cái"
        )
        self.period = (
            timezone.localdate().replace(day=1) - timedelta(days=1)
        ).strftime("%Y-%m")
        self.start, self.end = reports.period_dates(self.period)
        self.doc = Document.objects.create(
            kind="TRANSFER",
            status="TRANSIT",
            source=self.a,
            destination=self.b,
            created_by=self.admin,
            note="=1+1",
        )
        Line.objects.create(
            document=self.doc, product=self.p, quantity=2, unit_price=10
        )
        self.movement = Movement.objects.create(
            document=self.doc,
            warehouse=self.a,
            product=self.p,
            delta=-2,
            value_delta=-20,
            balance_after=3,
            value_after=30,
        )
        Movement.objects.filter(pk=self.movement.pk).update(created_at=self.start)
        # Boundary at next month belongs to the next period and must be excluded.
        boundary = Movement.objects.create(
            document=self.doc,
            warehouse=self.b,
            product=self.p,
            delta=2,
            value_delta=20,
            balance_after=2,
            value_after=20,
        )
        Movement.objects.filter(pk=boundary.pk).update(created_at=self.end)
        self.balance = Balance.objects.create(
            warehouse=self.a, product=self.p, quantity=3, value=30
        )
        self.configured = DriveConnection.objects.create(
            account_email="owner@example.test",
            encrypted_refresh_token="not-returned",
            connected_by=self.admin,
            folder_id=settings.GOOGLE_DRIVE_EXPECTED_FOLDER_ID,
        )

    def payload(self):
        return reports.build_report(self.period)

    def sheet_metadata(self, summary, state="pending"):
        return {
            "id": "synthetic-sheet",
            "mimeType": reports.SHEET_MIME,
            "appProperties": {
                "dhd_period": self.period,
                "dhd_sha256": summary["sha256"],
                "dhd_rows": summary["rows"],
                "dhd_state": state,
            },
        }

    def test_report_uses_executed_legs_local_period_boundaries_and_preserves_source(
        self,
    ):
        payload, summary = self.payload()
        book = load_workbook(BytesIO(payload), data_only=False)
        journal = list(book["Nhat ky kho"].values)
        self.assertEqual(len(journal), 2)
        self.assertEqual(journal[1][4], "A")
        self.assertEqual(journal[1][5], "00123")
        self.assertEqual(journal[1][8], "-2.000")
        self.assertEqual(journal[1][9], "-20.00")
        self.assertEqual(book["Nhat ky kho"]["G2"].data_type, "s")
        self.assertEqual(book["Phieu lien quan"]["I2"].data_type, "s")
        self.assertIn("Tồn, mượn", book["Thong tin"]["A7"].value)
        self.assertNotIn("password", book.sheetnames)
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.quantity, 3)
        self.assertEqual(Movement.objects.count(), 2)
        with patch.object(
            reports.drive,
            "google_request",
            side_effect=[payload, self.sheet_metadata(summary, "verified")],
        ):
            self.assertEqual(
                reports.verify_sheet(self.sheet_metadata(summary), "test-token")[
                    "state"
                ],
                "verified",
            )

    @patch("inventory.storage_reports.drive.google_request")
    def test_invalid_and_unfinished_period_never_contacts_google(self, google):
        for period in (
            "",
            "abc",
            "2026-13",
            "2026-00",
            "9999-12",
            timezone.localdate().strftime("%Y-%m"),
            "2026-01' or 1=1",
        ):
            with self.subTest(period=period), self.assertRaises(ValidationError):
                reports.export_month(period)
        google.assert_not_called()

    def test_row_and_file_limits_abort_instead_of_truncating(self):
        with patch.object(reports, "MAX_ROWS", 1), self.assertRaises(ValidationError):
            self.payload()
        with patch.object(reports, "MAX_BYTES", 10), self.assertRaises(ValidationError):
            self.payload()
        self.assertEqual(Movement.objects.count(), 2)

    def test_edited_sheet_fails_before_verified_metadata_and_can_be_retried(self):
        payload, summary = self.payload()
        book = load_workbook(BytesIO(payload))
        book["Nhat ky kho"]["I2"] = "999"
        buffer = BytesIO()
        book.save(buffer)
        with patch.object(
            reports.drive, "google_request", return_value=buffer.getvalue()
        ) as google:
            with self.assertRaisesMessage(ValidationError, "không khớp"):
                reports.verify_sheet(
                    self.sheet_metadata(summary, "verified"), "test-token"
                )
            self.assertEqual(google.call_count, 1)
        book["Nhat ky kho"]["I2"] = "=1+1"
        buffer = BytesIO()
        book.save(buffer)
        with (
            patch.object(
                reports.drive, "google_request", return_value=buffer.getvalue()
            ),
            self.assertRaises(ValidationError),
        ):
            reports.verify_sheet(self.sheet_metadata(summary), "test-token")
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.quantity, 3)

    def test_long_content_is_rejected_instead_of_silently_cut_by_excel(self):
        Document.objects.filter(pk=self.doc.pk).update(note="a" * 33000)
        with self.assertRaisesMessage(ValidationError, "quá dài"):
            self.payload()
        self.assertEqual(Movement.objects.count(), 2)

    @patch("inventory.storage_reports.drive.check_connection")
    @patch(
        "inventory.storage_reports.drive.folder",
        side_effect=["reports-root", "reports-year"],
    )
    @patch("inventory.storage_reports.drive.access_token", return_value="test-token")
    @patch("inventory.storage_reports.require_private_folder")
    def test_retry_verifies_existing_without_creating_another_sheet(
        self, private_folder, _token, folder, check
    ):
        payload, summary = self.payload()
        metadata = self.sheet_metadata(summary)
        with patch.object(
            reports.drive,
            "google_request",
            side_effect=[
                {"files": [metadata]},
                payload,
                self.sheet_metadata(summary, "verified"),
            ],
        ) as google:
            result = reports.export_month(self.period)
        self.assertEqual(result["state"], "verified")
        self.assertEqual(
            [call.kwargs.get("method", "GET") for call in google.call_args_list],
            ["GET", "GET", "PATCH"],
        )
        self.assertEqual(
            folder.call_args_list[0].args[1], settings.GOOGLE_DRIVE_EXPECTED_FOLDER_ID
        )
        check.assert_called_once()

    @patch("inventory.storage_reports.drive.check_connection")
    @patch(
        "inventory.storage_reports.drive.folder",
        side_effect=["reports-root", "reports-year"],
    )
    @patch("inventory.storage_reports.drive.access_token", return_value="test-token")
    def test_creation_is_in_connected_folder_never_changes_sharing_and_round_trips(
        self, _token, _folder, _check
    ):
        uploaded = {}

        def google(url, **kwargs):
            if "/permissions?" in url:
                return {"permissions": [{"type": "user"}]}
            if url.startswith(API + "?"):
                return {"files": []}
            if "upload/drive/v3/files" in url:
                body = kwargs["payload"]
                boundary = kwargs["content_type"].split("boundary=")[1].encode()
                parts = body.split(b"--" + boundary)
                meta = json.loads(parts[1].split(b"\r\n\r\n", 1)[1].strip())
                uploaded["bytes"] = parts[2].split(b"\r\n\r\n", 1)[1][:-2]
                uploaded["metadata"] = {"id": "synthetic-sheet", **meta}
                self.assertEqual(meta["parents"], ["reports-year"])
                self.assertNotIn("permissions", meta)
                return uploaded["metadata"]
            if "/export?" in url:
                return uploaded["bytes"]
            self.assertEqual(kwargs["method"], "PATCH")
            self.assertNotIn("permissions", url)
            return uploaded["metadata"]

        with patch.object(reports.drive, "google_request", side_effect=google):
            result = reports.export_month(self.period)
        self.assertEqual(result["state"], "verified")
        self.assertEqual(
            result["url"], "https://docs.google.com/spreadsheets/d/synthetic-sheet/edit"
        )

    @patch("inventory.storage_reports.drive.access_token")
    def test_wrong_root_fails_before_google_or_export(self, token):
        self.configured.folder_id = "other-folder"
        self.configured.save()
        with self.assertRaisesMessage(ValidationError, "khác thư mục"):
            reports.export_month(self.period)
        token.assert_not_called()

    def test_public_permission_in_later_page_is_rejected_without_changing_acl(self):
        with patch.object(
            reports.drive,
            "google_request",
            side_effect=[
                {"permissions": [{"type": "user"}], "nextPageToken": "next-page"},
                {"permissions": [{"type": "anyone"}]},
            ],
        ) as google:
            with self.assertRaisesMessage(ValidationError, "bất kỳ ai"):
                reports.require_private_folder("test-token", "reports-root")
            self.assertEqual(google.call_count, 2)
            self.assertTrue(
                all(
                    call.kwargs.get("method", "GET") == "GET"
                    for call in google.call_args_list
                )
            )
        with (
            patch.object(reports.drive, "google_request", return_value={}),
            self.assertRaises(ValidationError),
        ):
            reports.require_private_folder("test-token", "reports-root")

    @skipUnless(connection.vendor == "postgresql", "PostgreSQL report consistency")
    def test_concurrent_warehouse_change_is_not_blocked_and_report_uses_one_snapshot(
        self,
    ):
        changed = False

        def concurrent_change(execute, statement, params, many, context):
            nonlocal changed
            if (
                not changed
                and statement.startswith("SELECT")
                and 'FROM "inventory_balance"' in statement
            ):
                changed = True
                with reports.direct_connection() as db:
                    db.execute("SET statement_timeout='2s'")
                    db.execute(
                        "UPDATE inventory_balance SET quantity=4,value=40 WHERE id=%s",
                        (self.balance.pk,),
                    )
            return execute(statement, params, many, context)

        with connection.execute_wrapper(concurrent_change):
            payload, _summary = self.payload()
        self.assertTrue(changed)
        book = load_workbook(BytesIO(payload))
        self.assertEqual(book["Ton tai thoi diem xuat"]["E2"].value, "3.000")
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.quantity, 4)

    @skipUnless(connection.vendor == "postgresql", "Cross-process export locking")
    def test_two_workers_cannot_export_same_month_simultaneously(self):
        with reports.export_lock(self.period):
            with self.assertRaisesMessage(ValidationError, "Đang xuất báo cáo"):
                with reports.export_lock(self.period):
                    self.fail("Second worker obtained the same period lock")
        # Failure/exit releases the session lock; a retry can proceed.
        with reports.export_lock(self.period):
            pass

    @skipUnless(connection.vendor == "postgresql", "PostgreSQL storage metrics")
    @override_settings(DATABASE_STORAGE_LIMIT_MB=1024)
    def test_real_pg_size_accounts_for_snapshots_blobs_and_indexes_without_changing_data(
        self,
    ):
        schema = "dhd_backup_inventory_storage_test"
        with connection.cursor() as cursor:
            cursor.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(schema)))
            cursor.execute(
                sql.SQL(
                    "CREATE TABLE {}.sample AS SELECT repeat('test',1024) AS data"
                ).format(sql.Identifier(schema))
            )
        try:
            StoredFile.objects.create(
                sha256="a" * 64,
                context_key="storage-test",
                content_type="application/pdf",
                size=7,
                data=b"testpdf",
            )
            result = measure_storage()
            self.assertGreater(result["database_bytes"], 0)
            self.assertGreater(result["snapshot_bytes"], 0)
            self.assertEqual(result["blob_bytes"], 7)
            self.assertEqual(
                result["remaining_bytes"],
                max(1024 * 1024 * 1024 - result["database_bytes"], 0),
            )
            self.assertTrue(
                any(t["table"] == "inventory_balance" for t in result["tables"])
            )
            out = StringIO()
            call_command("storage_audit", stdout=out)
            self.assertNotIn("testpdf", out.getvalue())
            self.balance.refresh_from_db()
            self.assertEqual(self.balance.quantity, 3)
        finally:
            with connection.cursor() as cursor:
                cursor.execute(
                    sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(schema))
                )
