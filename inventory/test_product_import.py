import re
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from io import BytesIO
from threading import Barrier
from unittest import skipUnless
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections, connection, connections
from django.test import TestCase, TransactionTestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import Balance, CustomField, Document, Movement, Product, ProductPhoto
from .product_import import (
    MAX_EXPANDED_BYTES,
    MAX_PRODUCTS,
    MAX_SCAN_ROWS,
    MAX_UPLOAD_BYTES,
    confirm_product_import,
    parse_product_workbook,
)
from .test_support import grant_functions

HEADERS = ["Mã hàng", "Tên hàng", "Đơn vị tính"]
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_upload(rows=(), headers=None, *, filename="hang-hoa.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hàng hóa"
    sheet.append(HEADERS if headers is None else headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(filename, output.getvalue(), content_type=XLSX_TYPE)


class ProductImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.creator = User.objects.create_user("product-importer", password="1")
        cls.viewer = User.objects.create_user("product-viewer", password="1")
        cls.receipt_importer = User.objects.create_user(
            "receipt-importer", password="1"
        )
        cls.admin = User.objects.create_superuser(
            "product-admin", "admin@example.test", "1"
        )
        grant_functions(cls.creator, {"products_create", "products_view"})
        grant_functions(cls.viewer, {"products_view"})
        grant_functions(cls.receipt_importer, {"tools_import"})

    def parse(self, rows=(), headers=None, **kwargs):
        return parse_product_workbook(workbook_upload(rows, headers, **kwargs))

    def confirm(self, preview, *, user=None):
        return confirm_product_import(
            user or self.creator,
            preview["rows"],
            preview["schema"],
            filename=preview["filename"],
        )

    def preview_http(self, rows=None, headers=None):
        self.client.force_login(self.creator)
        response = self.client.post(
            reverse("product_import"),
            {
                "action": "preview",
                "file": workbook_upload(
                    rows or [["NEW-001", "Máy kiểm kho", "Cái"]], headers
                ),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("product_import", self.client.session)
        return response, self.client.session["product_import"]["key"]

    def assert_no_stock_writes(self):
        self.assertFalse(Balance.objects.exists())
        self.assertFalse(Document.objects.exists())
        self.assertFalse(Movement.objects.exists())

    def test_real_downloaded_template_round_trip_preserves_text_sku_and_qr(self):
        self.client.force_login(self.creator)
        response = self.client.get(reverse("product_import_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        payload = (
            b"".join(response.streaming_content)
            if response.streaming
            else response.content
        )
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook.worksheets[0]
        for column, value in enumerate(["000042", "Bút thử điện", "Cái"], 1):
            sheet.cell(2, column, value)
        output = BytesIO()
        workbook.save(output)
        preview = parse_product_workbook(
            SimpleUploadedFile("mau-da-dien.xlsx", output.getvalue())
        )
        self.assertTrue(preview["valid"], preview)
        created = self.confirm(preview)
        self.assertEqual(len(created), 1)
        product = Product.objects.get(sku="000042")
        self.assertEqual(product.name, "Bút thử điện")
        self.assertIsNotNone(product.qr_token)
        self.assertFalse(ProductPhoto.objects.exists())
        self.assert_no_stock_writes()
        self.client.force_login(self.admin)
        qr = self.client.get(reverse("product_qr", args=[product.pk]))
        self.assertEqual(qr.status_code, 200)
        self.assertEqual(qr["Content-Type"], "image/png")
        self.assertTrue(qr.content.startswith(b"\x89PNG"))

    def test_all_optional_and_custom_fields_are_saved_without_drive_connection(self):
        CustomField.objects.create(key="origin", label="Xuất xứ")
        keys = [
            "category",
            "brand",
            "model_number",
            "specifications",
            "storage_instructions",
            "minimum",
        ]
        headers = HEADERS + [Product._meta.get_field(key).verbose_name for key in keys]
        headers += ["Xuất xứ"]
        preview = self.parse(
            [
                [
                    "IMG-001",
                    "Cân điện tử",
                    "Bộ",
                    "Thiết bị",
                    "Lemon",
                    "L-1",
                    "20 kg",
                    "Khô ráo",
                    "1.250",
                    "Việt Nam",
                ]
            ],
            headers,
        )
        self.assertTrue(preview["valid"], preview)
        self.confirm(preview)
        product = Product.objects.get(sku="IMG-001")
        self.assertEqual(str(product.minimum), "1.250")
        self.assertEqual(product.extra, {"origin": "Việt Nam"})
        self.assertEqual(product.model_number, "L-1")
        self.assertEqual(product.specifications, "20 kg")
        self.assertEqual(product.storage_instructions, "Khô ráo")
        self.assertFalse(ProductPhoto.objects.exists())
        self.assert_no_stock_writes()

    def test_template_distinguishes_custom_label_that_matches_builtin_field(self):
        CustomField.objects.create(key="legacy_sku", label="Mã hàng")
        self.client.force_login(self.creator)
        response = self.client.get(reverse("product_import_template"))
        payload = (
            b"".join(response.streaming_content)
            if response.streaming
            else response.content
        )
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook.worksheets[0]
        values = {
            "Mã hàng": "CURRENT-001",
            "Tên hàng": "Tên hiện tại",
            "Đơn vị tính": "Cái",
            "Bổ sung: Mã hàng": "LEGACY-009",
        }
        for cell in sheet[1]:
            if cell.value in values:
                sheet.cell(2, cell.column, values[cell.value])
        output = BytesIO()
        workbook.save(output)
        preview = parse_product_workbook(
            SimpleUploadedFile("legacy.xlsx", output.getvalue())
        )
        self.assertTrue(preview["valid"], preview)
        self.confirm(preview)
        product = Product.objects.get(sku="CURRENT-001")
        self.assertEqual(product.extra["legacy_sku"], "LEGACY-009")

    def test_empty_optional_fields_default_to_zero_and_each_product_gets_own_qr(self):
        preview = self.parse(
            [["QR-001", "Hàng một", "Cái"], ["QR-002", "Hàng hai", "Bộ"]]
        )
        self.assertTrue(preview["valid"], preview)
        self.confirm(preview)
        self.assertEqual(Product.objects.count(), 2)
        self.assertEqual(Product.objects.values("qr_token").distinct().count(), 2)
        self.assertTrue(all(product.minimum == 0 for product in Product.objects.all()))
        self.assert_no_stock_writes()

    def test_preview_makes_no_database_changes_and_reports_excel_row_numbers(self):
        preview = self.parse(
            [["GOOD-001", "Hàng hợp lệ", "Cái"], ["", "Thiếu mã", "Bộ"]]
        )
        self.assertFalse(preview["valid"])
        self.assertEqual(preview["count"], 2)
        self.assertEqual(preview["valid_count"], 1)
        self.assertEqual([row["row_number"] for row in preview["rows"]], [2, 3])
        self.assertTrue(preview["rows"][1]["errors"])
        self.assertFalse(Product.objects.exists())
        with self.assertRaises(ValidationError):
            self.confirm(preview)
        self.assertFalse(Product.objects.exists())

    def test_required_field_length_and_minimum_validation_blocks_whole_batch(self):
        headers = HEADERS + [Product._meta.get_field("minimum").verbose_name]
        invalid_rows = [
            ["", "Thiếu mã", "Cái", "0"],
            ["NO-NAME", "", "Cái", "0"],
            ["NO-UNIT", "Tên hàng", "", "0"],
            ["X" * 51, "Mã quá dài", "Cái", "0"],
            ["LONG-NAME", "X" * 201, "Cái", "0"],
            ["LONG-UNIT", "Đơn vị quá dài", "X" * 31, "0"],
            ["NEGATIVE", "Tồn tối thiểu âm", "Cái", "-1"],
            ["PRECISION", "Quá ba số thập phân", "Cái", "1.0001"],
            ["NAN", "Giá trị không hữu hạn", "Cái", "NaN"],
        ]
        for invalid in invalid_rows:
            with self.subTest(sku=invalid[0]):
                preview = self.parse(
                    [["VALID", "Hàng đúng", "Cái", "0"], invalid], headers
                )
                self.assertFalse(preview["valid"])
                with self.assertRaises(ValidationError):
                    self.confirm(preview)
                self.assertFalse(Product.objects.exists())

    def test_duplicate_file_skus_and_existing_skus_do_not_update_or_partially_add(self):
        existing = Product.objects.create(sku="EXISTING", name="Giữ nguyên", unit="Cái")
        cases = [
            [["DUP", "Một", "Cái"], [" DUP ", "Hai", "Bộ"]],
            [["NEW", "Hàng mới", "Cái"], ["EXISTING", "Không ghi đè", "Bộ"]],
        ]
        for rows in cases:
            with self.subTest(rows=rows):
                preview = self.parse(rows)
                self.assertFalse(preview["valid"])
                with self.assertRaises(ValidationError):
                    self.confirm(preview)
                self.assertEqual(Product.objects.count(), 1)
                existing.refresh_from_db()
                self.assertEqual(existing.name, "Giữ nguyên")
                self.assertEqual(existing.unit, "Cái")

    def test_formulas_error_cells_and_numeric_skus_are_not_silently_imported(self):
        for row in [
            [123, "Mã Excel kiểu số", "Cái"],
            ["FORMULA", '=CONCAT("Hàng"," mới")', "Cái"],
            ["ERROR", "#DIV/0!", "Cái"],
        ]:
            with self.subTest(row=row):
                preview = self.parse([row])
                self.assertFalse(preview["valid"])
                self.assertTrue(preview["errors"] or preview["rows"][0]["errors"])
                with self.assertRaises(ValidationError):
                    self.confirm(preview)
        self.assertFalse(Product.objects.exists())

    def test_unknown_or_repeated_headers_and_missing_required_columns_are_rejected(
        self,
    ):
        cases = [
            (HEADERS + ["Số tồn ban đầu"], [["A", "Hàng A", "Cái", 9]]),
            (HEADERS + ["Mã hàng"], [["A", "Hàng A", "Cái", "B"]]),
            (["Mã hàng", "Tên hàng"], [["A", "Hàng A"]]),
        ]
        for headers, rows in cases:
            with self.subTest(headers=headers):
                preview = self.parse(rows, headers)
                self.assertFalse(preview["valid"])
                self.assertTrue(preview["errors"])
        self.assertFalse(Product.objects.exists())

    def test_empty_corrupt_legacy_and_oversized_files_are_rejected(self):
        cases = [
            workbook_upload(),
            workbook_upload([["A", "Hàng A", "Cái"]], filename="cu.xls"),
            SimpleUploadedFile("hong.xlsx", b"not an Excel workbook"),
            SimpleUploadedFile("qua-lon.xlsx", b"x" * (MAX_UPLOAD_BYTES + 1)),
        ]
        for upload in cases:
            with self.subTest(filename=upload.name, size=upload.size):
                preview = parse_product_workbook(upload)
                self.assertFalse(preview["valid"])
                self.assertTrue(preview["errors"])
        self.assertFalse(Product.objects.exists())

    def test_workbook_expanded_zip_limit_prevents_decompression_bomb(self):
        upload = workbook_upload([["A", "Hàng A", "Cái"]])
        source = BytesIO(upload.read())
        output = BytesIO()
        with ZipFile(source) as original, ZipFile(output, "w", ZIP_DEFLATED) as changed:
            for item in original.infolist():
                changed.writestr(item, original.read(item.filename))
            changed.writestr(
                "oversized-unused-part.xml", b"x" * (MAX_EXPANDED_BYTES + 1)
            )
        preview = parse_product_workbook(
            SimpleUploadedFile("zip-lon.xlsx", output.getvalue())
        )
        self.assertFalse(preview["valid"])
        self.assertTrue(preview["errors"])
        self.assertFalse(Product.objects.exists())

    def test_batch_row_cap_is_enforced_before_any_product_is_created(self):
        preview = self.parse(
            [
                [f"BULK-{index}", f"Hàng {index}", "Cái"]
                for index in range(MAX_PRODUCTS + 1)
            ]
        )
        self.assertFalse(preview["valid"])
        self.assertTrue(preview["errors"])
        self.assertFalse(Product.objects.exists())

    def test_maximum_valid_batch_uses_batch_validation_and_creates_all_qr_codes(self):
        CustomField.objects.create(key="serial_group", label="Nhóm số hiệu")
        with CaptureQueriesContext(connection) as queries:
            preview = self.parse(
                [
                    [f"MAX-{index:04d}", f"Hàng {index}", "Cái"]
                    for index in range(MAX_PRODUCTS)
                ]
            )
        self.assertTrue(preview["valid"], preview["errors"])
        self.assertLessEqual(
            len(queries), 10, "Không kiểm tra SKU bằng một truy vấn cho mỗi dòng"
        )
        products = self.confirm(preview)
        self.assertEqual(len(products), MAX_PRODUCTS)
        self.assertEqual(Product.objects.count(), MAX_PRODUCTS)
        self.assertEqual(
            Product.objects.values("qr_token").distinct().count(), MAX_PRODUCTS
        )
        self.assert_no_stock_writes()

    def test_forged_dimensions_cannot_hide_cells_beyond_the_scanning_limit(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(HEADERS)
        sheet.append(["A", "Hàng A", "Cái"])
        sheet.cell(MAX_SCAN_ROWS + 1, 1, "HIDDEN")
        raw = BytesIO()
        workbook.save(raw)
        output = BytesIO()
        with (
            ZipFile(BytesIO(raw.getvalue())) as original,
            ZipFile(output, "w", ZIP_DEFLATED) as changed,
        ):
            for item in original.infolist():
                content = original.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(
                        rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', content
                    )
                changed.writestr(item, content)
        preview = parse_product_workbook(
            SimpleUploadedFile("hidden.xlsx", output.getvalue())
        )
        self.assertFalse(preview["valid"])
        self.assertTrue(preview["errors"])
        self.assertFalse(Product.objects.exists())

    def test_stale_products_and_custom_field_schema_are_rechecked_on_confirmation(self):
        preview = self.parse(
            [["AFTER-PREVIEW", "Hàng chờ", "Cái"], ["STILL-NEW", "Hàng còn mới", "Cái"]]
        )
        self.assertTrue(preview["valid"], preview)
        Product.objects.create(sku="AFTER-PREVIEW", name="Người khác đã tạo", unit="Bộ")
        with self.assertRaises(ValidationError):
            self.confirm(preview)
        self.assertFalse(Product.objects.filter(sku="STILL-NEW").exists())
        clean_preview = self.parse([["SCHEMA", "Chờ cấu hình", "Cái"]])
        CustomField.objects.create(key="warranty", label="Thời hạn bảo hành")
        with self.assertRaises(ValidationError):
            self.confirm(clean_preview)
        self.assertFalse(Product.objects.filter(sku="SCHEMA").exists())

    def test_confirmation_revalidates_rows_and_ignores_claimed_preview_validity(self):
        preview = self.parse([["SAFE", "Hàng đúng", "Cái"]])
        invalid = deepcopy(preview)
        invalid["rows"][0]["data"]["minimum"] = "-1"
        invalid["rows"][0]["errors"] = []
        with self.assertRaises(ValidationError):
            self.confirm(invalid)
        tampered = deepcopy(preview)
        tampered["rows"][0]["data"]["quantity"] = "999"
        with self.assertRaises(ValidationError):
            self.confirm(tampered)
        self.assertFalse(Product.objects.exists())
        self.assert_no_stock_writes()

    def test_routes_and_template_use_product_creation_permission_not_receipt_import(
        self,
    ):
        for route in ["product_import", "product_import_template"]:
            url = reverse(route)
            self.assertEqual(self.client.get(url).status_code, 302)
            for user in [self.viewer, self.receipt_importer]:
                self.client.force_login(user)
                self.assertEqual(self.client.get(url).status_code, 403)
            for user in [self.creator, self.admin]:
                self.client.force_login(user)
                self.assertEqual(self.client.get(url).status_code, 200)
            self.client.logout()

    def test_preview_confirm_replay_and_forged_post_fields_do_not_duplicate_or_override(
        self,
    ):
        _, key = self.preview_http()
        self.assertFalse(Product.objects.exists())
        response = self.client.post(
            reverse("product_import"),
            {
                "action": "confirm",
                "import_key": key,
                "sku": "HACK",
                "minimum": "999",
                "quantity": "999",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Product.objects.count(), 1)
        product = Product.objects.get()
        self.assertEqual(product.sku, "NEW-001")
        self.assertEqual(product.minimum, 0)
        self.assertNotIn("product_import", self.client.session)
        replay = self.client.post(
            reverse("product_import"), {"action": "confirm", "import_key": key}
        )
        self.assertIn(replay.status_code, [200, 302])
        self.assertEqual(Product.objects.count(), 1)
        self.assert_no_stock_writes()

    def test_cancel_forged_missing_and_replaced_preview_keys_cannot_import(self):
        _, first = self.preview_http([["FIRST", "Phiên cũ", "Cái"]])
        _, second = self.preview_http([["SECOND", "Phiên mới", "Cái"]])
        self.assertNotEqual(first, second)
        for key in [None, "not-a-preview-key", first]:
            payload = {"action": "confirm"}
            if key is not None:
                payload["import_key"] = key
            response = self.client.post(reverse("product_import"), payload)
            self.assertIn(response.status_code, [200, 302])
            self.assertFalse(Product.objects.exists())
        self.client.post(reverse("product_import"), {"action": "cancel"})
        self.assertEqual(self.client.session["product_import"]["key"], second)
        response = self.client.post(
            reverse("product_import"),
            {"action": "cancel", "import_key": second},
        )
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("product_import", self.client.session)
        self.client.post(
            reverse("product_import"), {"action": "confirm", "import_key": second}
        )
        self.assertFalse(Product.objects.exists())

    def test_http_confirmation_conflict_is_reported_without_partial_writes(self):
        _, key = self.preview_http(
            [["CONFLICT", "Hàng xem trước", "Cái"], ["UNTOUCHED", "Chưa tạo", "Cái"]]
        )
        Product.objects.create(sku="CONFLICT", name="Đã tạo ở cửa sổ khác", unit="Bộ")
        response = self.client.post(
            reverse("product_import"), {"action": "confirm", "import_key": key}
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Dòng 2")
        self.assertContains(response, "Mã hàng đã có")
        self.assertFalse(Product.objects.filter(sku="UNTOUCHED").exists())
        self.assertEqual(Product.objects.get(sku="CONFLICT").unit, "Bộ")
        self.assertNotIn("product_import", self.client.session)

    def test_expired_preview_cannot_be_confirmed(self):
        _, key = self.preview_http()
        preview = self.client.session["product_import"]
        with patch(
            "inventory.product_import_views.time.time",
            return_value=preview["created_at"] + 1801,
        ):
            response = self.client.post(
                reverse("product_import"), {"action": "confirm", "import_key": key}
            )
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("product_import", self.client.session)
        self.assertFalse(Product.objects.exists())

    def test_revoked_permission_after_preview_blocks_service_and_http_confirmation(
        self,
    ):
        _, key = self.preview_http()
        preview = self.client.session["product_import"]
        grant_functions(self.creator, {"products_view"})
        response = self.client.post(
            reverse("product_import"), {"action": "confirm", "import_key": key}
        )
        self.assertEqual(response.status_code, 403)
        self.creator.refresh_from_db()
        with self.assertRaises(PermissionDenied):
            self.confirm(preview)
        self.assertFalse(Product.objects.exists())


@skipUnless(
    connection.vendor == "postgresql", "Requires PostgreSQL concurrent transactions"
)
class ProductImportConcurrencyTests(TransactionTestCase):
    def test_overlapping_imports_from_different_users_are_atomic(self):
        User = get_user_model()
        users = [
            User.objects.create_user(f"concurrent-import-{index}") for index in range(2)
        ]
        previews = []
        for index, user in enumerate(users):
            grant_functions(user, {"products_create"})
            preview = parse_product_workbook(
                workbook_upload(
                    [
                        [f"UNIQUE-{index}", f"Hàng riêng {index}", "Cái"],
                        ["SHARED", "Hàng chung", "Cái"],
                    ]
                )
            )
            self.assertTrue(preview["valid"], preview)
            previews.append(preview)
        barrier = Barrier(2)

        def create(index):
            close_old_connections()
            try:
                user = User.objects.get(pk=users[index].pk)
                preview = previews[index]
                barrier.wait(timeout=10)
                confirm_product_import(user, preview["rows"], preview["schema"])
                return "created"
            except ValidationError:
                return "blocked"
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(create, [0, 1]))
        self.assertEqual(sorted(results), ["blocked", "created"])
        self.assertEqual(Product.objects.count(), 2)
        self.assertEqual(Product.objects.filter(sku="SHARED").count(), 1)
        self.assertEqual(Product.objects.filter(sku__startswith="UNIQUE-").count(), 1)
        self.assertFalse(Balance.objects.exists())
        self.assertFalse(Document.objects.exists())
        self.assertFalse(Movement.objects.exists())
