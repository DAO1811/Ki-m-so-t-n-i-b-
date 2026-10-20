"""Bounded product import with optional, separately approved receipt drafts."""

import hashlib
import json
import re
import unicodedata
from io import BytesIO
from zipfile import BadZipFile, LargeZipFile, ZipFile

from defusedxml.ElementTree import iterparse
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from .access import has_function, require_function
from .forms import ProductForm
from .models import CustomField, Product
from .product_allocation import (
    ALLOCATION_COLUMNS,
    ALLOCATION_KEYS,
    allocation_summary,
    create_allocated_receipts,
    has_allocation,
    validate_allocations,
)

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_PRODUCTS = 500
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_SCAN_ROWS = 2000
MAX_COLUMNS = 100
MAX_NORMALIZED_BYTES = 1024 * 1024
MAX_ZIP_ENTRIES = 200

BASE_COLUMNS = (
    (
        "sku",
        "Mã hàng",
        ("sku", "ma hang", "ma hang hoa", "ma san pham", "product code"),
    ),
    (
        "name",
        "Tên hàng",
        ("name", "ten hang", "ten hang hoa", "ten san pham", "product name"),
    ),
    ("unit", "Đơn vị tính", ("unit", "don vi", "don vi tinh", "dvt")),
    ("category", "Nhóm hàng", ("category", "nhom hang", "loai hang")),
    (
        "brand",
        "Nhãn hiệu / nhà sản xuất",
        ("brand", "nhan hieu", "thuong hieu", "nha san xuat"),
    ),
    (
        "model_number",
        "Model / mã nhà sản xuất",
        ("model_number", "model", "ma nha san xuat"),
    ),
    (
        "specifications",
        "Quy cách / thông số kỹ thuật",
        ("specifications", "quy cach", "thong so ky thuat"),
    ),
    (
        "storage_instructions",
        "Điều kiện bảo quản",
        ("storage_instructions", "dieu kien bao quan", "bao quan"),
    ),
    (
        "minimum",
        "Mức tồn tối thiểu mỗi kho",
        ("minimum", "ton toi thieu", "muc ton toi thieu", "min_stock"),
    ),
    ("qr_auto", "QR tự động", ("qr_auto", "tu dong qr")),
    ("qr_code", "Mã QR", ("qr_code", "noi dung qr", "ma qr")),
)
REQUIRED_COLUMNS = {"sku", "name", "unit"}


class ProductImportForm(ProductForm):
    validate_identity = False

    def validate_unique(self):
        # One batch query below replaces one uniqueness round-trip per product.
        pass


def _custom_fields():
    return list(CustomField.objects.order_by("pk"))


def _schema(fields):
    value = {"version": 2, "fields": [(field.key, field.label) for field in fields]}
    return hashlib.sha256(json.dumps(value, ensure_ascii=False).encode()).hexdigest()


def _columns(fields):
    return [
        {"key": key, "label": label}
        for key, label, _aliases in (*BASE_COLUMNS, *ALLOCATION_COLUMNS)
    ] + [
        {"key": "extra_" + field.key, "label": "Bổ sung: " + field.label}
        for field in fields
    ]


def _header(value):
    value = str(value or "").strip().lower().replace("đ", "d")
    value = "".join(
        char
        for char in unicodedata.normalize("NFD", value)
        if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def _aliases(fields):
    aliases = {}

    def add(alias, key):
        normalized = _header(alias)
        if normalized:
            aliases.setdefault(normalized, set()).add(key)

    for key, label, variants in (*BASE_COLUMNS, *ALLOCATION_COLUMNS):
        for alias in (key, label, *variants):
            add(alias, key)
    built_in_aliases = set(aliases)
    for field in fields:
        key = "extra_" + field.key
        for alias in (key, "Bổ sung: " + field.label):
            add(alias, key)
        if _header(field.label) not in built_in_aliases:
            add(field.label, key)
    return aliases


def build_product_template(*, user=None):
    """Return an XLSX stream with current custom columns and text-formatted SKU."""
    fields = _custom_fields()
    columns = _columns(fields)
    if len(columns) > MAX_COLUMNS:
        raise ValidationError(
            "Có quá nhiều trường bổ sung để tạo mẫu Excel (tối đa 100 cột)."
        )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hàng hóa"
    sheet.freeze_panes = "A2"
    sheet.append([column["label"] for column in columns])
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for index, column in enumerate(columns, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="116B5B")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = (
            40 if column["key"] in {"specifications", "storage_instructions"} else 26
        )
        for row in range(2, MAX_PRODUCTS + 2):
            sheet.cell(row, index).number_format = (
                "0.###"
                if column["key"] in {"minimum", "receipt_quantity"}
                else "0.####"
                if column["key"] == "receipt_price"
                else "@"
            )
    sheet.row_dimensions[1].height = 32
    instructions = workbook.create_sheet("Hướng dẫn")
    instructions.column_dimensions["A"].width = 110
    for text in (
        "NHẬP DANH MỤC HÀNG HÓA BẰNG EXCEL",
        "1. Điền dữ liệu từ dòng 2 của trang Hàng hóa; giữ tiêu đề ở dòng 1.",
        "2. Bắt buộc: Mã hàng, Tên hàng, Đơn vị tính. Phân bổ nhiều kho: lặp mã ở các dòng khác kho, giữ nguyên thông tin hàng hóa và QR.",
        "3. Mã hàng phải là văn bản (Text) để giữ số 0 đầu mã. Không dùng công thức.",
        "4. Mức tồn tối thiểu: số không âm, tối đa 3 chữ số thập phân; bỏ trống tương đương 0.",
        "5. Tối đa 500 dòng hàng, tệp .xlsx không quá 5 MB. Dòng trống được bỏ qua.",
        "6. Chỉ thêm mã mới. Mã đã có: dùng chức năng Nhập kho → Lập phiếu nhập bằng Excel. Không lặp cùng mã tại cùng kho.",
        "7. Ba cột Mã kho nhập, Số lượng nhập, Đơn giá nhập cùng để trống: chỉ thêm danh mục. Có phân bổ: điền đủ cả ba, đơn giá miễn phí ghi 0.",
        "7a. Mỗi kho tự tạo phiếu nhập chờ duyệt, tối đa 50 mã/phiếu. Chưa cộng tồn trước khi được duyệt đúng quyền. Ngày nhập là hôm nay.",
        "7b. Số lượng nhập dương, tối đa 3 số lẻ. Đơn giá không âm, tối đa 4 số lẻ. Mã kho lấy tại trang Danh sách kho; không dùng tên kho.",
        "8. Cột Bổ sung là thông tin hàng hóa đã được quản trị viên cấu hình, có thể để trống.",
        "9. Tải tệp để xem trước, sửa mọi lỗi trong Excel rồi tải lại trước khi xác nhận.",
        "10. Ảnh hàng hóa và hóa đơn được thêm ở trang chi tiết/phiếu nhập sau khi tạo mã. Cần quyền Tạo phiếu nhập kho và quyền tại kho được chọn.",
        "11. QR tự động: bỏ trống hoặc Có. Nếu dùng QR sẵn/tự đặt, điền Không và nhập Mã QR. QR cố định suốt vòng đời hàng hóa.",
    ):
        instructions.append([text])
    warehouses = workbook.create_sheet("Danh sách kho")
    warehouses.append(["Mã kho nhập", "Tên kho"])
    warehouses.column_dimensions["A"].width = 28
    warehouses.column_dimensions["B"].width = 50
    if user is not None and has_function(user, "receipts_create"):
        from .services import warehouses_for

        for warehouse in warehouses_for(user).order_by("code"):
            warehouses.append([warehouse.code, warehouse.name])
            warehouses.cell(warehouses.max_row, 1).number_format = "@"
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _read_workbook(upload):
    filename = str(getattr(upload, "name", "hang-hoa.xlsx"))
    if not filename.lower().endswith(".xlsx"):
        raise ValidationError(
            "Chỉ nhận tệp Excel .xlsx. Hãy lưu lại tệp theo định dạng này."
        )
    size = getattr(upload, "size", None)
    if size is not None and size > MAX_UPLOAD_BYTES:
        raise ValidationError("Tệp Excel vượt quá 5 MB.")
    raw = upload.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValidationError("Tệp Excel vượt quá 5 MB.")
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if (
                len(entries) > MAX_ZIP_ENTRIES
                or len(names) != len(entries)
                or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES
                or any(entry.flag_bits & 1 for entry in entries)
                or any("vbaproject" in entry.filename.lower() for entry in entries)
            ):
                raise ValidationError(
                    "Cấu trúc tệp Excel quá lớn hoặc không được hỗ trợ. Hãy dùng mẫu .xlsx mới."
                )
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise ValidationError("Tệp không phải là bảng tính Excel .xlsx hợp lệ.")
            # Bound actual XML coordinates before openpyxl fills missing rows
            # or columns. A forged narrow dimension must not hide cell A999999.
            for entry in entries:
                if not (
                    entry.filename.startswith("xl/worksheets/")
                    and entry.filename.endswith(".xml")
                ):
                    continue
                with archive.open(entry) as stream:
                    for _event, element in iterparse(stream, events=("start",)):
                        tag = element.tag.rsplit("}", 1)[-1]
                        if tag == "row":
                            reference = element.attrib.get("r", "")
                            if reference and (
                                not reference.isascii()
                                or not reference.isdigit()
                                or len(reference) > 4
                                or int(reference) > MAX_SCAN_ROWS
                            ):
                                raise ValidationError(
                                    "Trang tính có dòng nằm ngoài vùng hỗ trợ. Hãy chép 500 mặt hàng vào mẫu mới."
                                )
                        elif tag == "c":
                            reference = element.attrib.get("r", "")
                            if reference:
                                match = re.fullmatch(
                                    r"([A-Z]{1,3})([1-9][0-9]{0,6})", reference
                                )
                                if (
                                    not match
                                    or column_index_from_string(match[1]) > MAX_COLUMNS
                                    or int(match[2]) > MAX_SCAN_ROWS
                                ):
                                    raise ValidationError(
                                        "Trang tính có ô nằm ngoài vùng hỗ trợ. Hãy chép dữ liệu vào mẫu mới (tối đa 100 cột)."
                                    )
                        element.clear()
        return load_workbook(
            BytesIO(raw), read_only=True, data_only=False, keep_links=False
        )
    except ValidationError:
        raise
    except (BadZipFile, LargeZipFile, OSError, ValueError, KeyError, TypeError) as exc:
        raise ValidationError(
            "Không đọc được tệp Excel. Tệp có thể bị hỏng hoặc không đúng định dạng .xlsx."
        ) from exc
    except Exception as exc:
        # XML/relationship parsing errors vary between openpyxl versions.
        raise ValidationError(
            "Không đọc được cấu trúc Excel. Hãy chép dữ liệu vào tệp mẫu mới."
        ) from exc


def _error(row, field, label, message):
    row["errors"].append({"field": field, "label": label, "message": str(message)})


def _validate_rows(rows, fields):
    labels = {column["key"]: column["label"] for column in _columns(fields)}
    seen = {}
    candidates = set()
    for row in rows:
        allocation = {key: row["data"].get(key, "") for key in ALLOCATION_KEYS}
        form = ProductImportForm(data=row["data"], custom_fields=fields)
        if form.is_valid():
            instance = form.save(commit=False)
            row["data"] = {
                key: str(form.cleaned_data.get(key, "") or "") for key in labels
            }
            # Decimal zero must remain explicit, and no ModelForm defaults may
            # accidentally fill the required unit of an incomplete source row.
            row["data"].update(allocation)
            row["data"]["minimum"] = str(form.cleaned_data["minimum"])
            row["data"]["qr_auto"] = "True" if form.cleaned_data["qr_auto"] else "False"
            row["instance"] = instance
        else:
            for key, errors in form.errors.items():
                for message in errors:
                    _error(row, key, labels.get(key, "Thông tin hàng"), message)
        sku = str(row["data"].get("sku", "")).strip()
        if sku:
            candidates.add(sku)
            seen.setdefault(sku, []).append(row)
    for sku, duplicates in seen.items():
        if len(duplicates) < 2 or not all(
            has_allocation(row) and row.get("instance") for row in duplicates
        ):
            continue

        def master_data(row):
            return {
                key: value
                for key, value in row["data"].items()
                if key not in ALLOCATION_KEYS
                and not (key == "qr_code" and row["data"].get("qr_auto") == "True")
            }

        if any(
            master_data(row) != master_data(duplicates[0]) for row in duplicates[1:]
        ):
            for row in duplicates:
                _error(
                    row,
                    "sku",
                    labels["sku"],
                    "Cùng mã phân bổ nhiều kho phải giữ nguyên tên, đơn vị, QR và thông tin hàng hóa.",
                )
        else:
            for row in duplicates:
                row["instance"] = duplicates[0]["instance"]
                row["data"]["qr_code"] = duplicates[0]["data"]["qr_code"]
    existing = set(
        Product.objects.filter(sku__in=candidates).values_list("sku", flat=True)
    )
    code_rows = {}
    for row in rows:
        instance = row.get("instance")
        if instance:
            code_rows.setdefault(instance.qr_code, []).append(row)
    codes = set(code_rows)
    reserved_codes = set(
        Product.objects.filter(qr_code__in=candidates | codes).values_list(
            "qr_code", flat=True
        )
    )
    occupied_skus = set(
        Product.objects.filter(sku__in=codes).values_list("sku", flat=True)
    )
    for row in rows:
        instance = row.get("instance")
        if not instance:
            continue
        code = instance.qr_code
        if (
            len({other["instance"].sku for other in code_rows[code]}) > 1
            or code in reserved_codes
            or code in occupied_skus
            or any(
                other["data"].get("sku") != instance.sku for other in seen.get(code, [])
            )
        ):
            _error(
                row,
                "qr_code",
                "Mã QR",
                "Mã QR trùng với QR hoặc mã hàng khác. Mỗi mã QR chỉ gắn với một hàng hóa.",
            )
        if instance.sku in reserved_codes:
            _error(
                row, "sku", "Mã hàng", "Mã hàng đã được dùng làm QR của hàng hóa khác."
            )
    for sku, duplicates in seen.items():
        for row in duplicates:
            allocation_repeat = all(
                has_allocation(other) for other in duplicates
            ) and len(
                {other["data"].get("warehouse_code") for other in duplicates}
            ) == len(duplicates)
            if len(duplicates) > 1 and not allocation_repeat:
                lines = ", ".join(
                    str(duplicate["row_number"]) for duplicate in duplicates[:10]
                )
                if len(duplicates) > 10:
                    lines += f" và {len(duplicates) - 10} dòng khác"
                _error(
                    row,
                    "sku",
                    labels["sku"],
                    f"Mã hàng trùng trong tệp tại dòng {lines}.",
                )
            if sku in existing:
                _error(
                    row,
                    "sku",
                    labels["sku"],
                    "Mã hàng đã có trong phần mềm. Chức năng này chỉ thêm hàng mới.",
                )
    return rows


def parse_product_workbook(upload, *, user=None):
    """Return a JSON-safe preview; invalid rows are never silently imported."""
    fields = _custom_fields()
    columns = _columns(fields)
    result = {
        "filename": str(getattr(upload, "name", "hang-hoa.xlsx"))[:250],
        "schema": _schema(fields),
        "business_date": str(timezone.localdate()),
        "columns": columns,
        "rows": [],
        "errors": [],
        "valid": False,
        "count": 0,
        "valid_count": 0,
    }
    workbook = None
    validated = False
    try:
        if len(columns) > MAX_COLUMNS:
            raise ValidationError(
                "Danh mục đang có quá nhiều trường bổ sung (tối đa 100 cột)."
            )
        workbook = _read_workbook(upload)
        if not workbook.worksheets or len(workbook.worksheets) > 10:
            raise ValidationError(
                "Tệp Excel cần có từ 1 đến 10 trang tính; dữ liệu hàng hóa nằm ở trang đầu tiên."
            )
        sheet = workbook.worksheets[0]
        if (sheet.max_row or 0) > MAX_SCAN_ROWS or (
            sheet.max_column or 0
        ) > MAX_COLUMNS:
            raise ValidationError(
                "Trang tính quá lớn. Chỉ giữ vùng dữ liệu tối đa 500 mặt hàng và 100 cột."
            )
        # Reset declared dimensions: an incorrect A1:A1 dimension must not hide
        # later rows from validation. Real row/column limits are checked below.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header_cells = next(iterator, ())
        if not header_cells:
            raise ValidationError("Tệp Excel chưa có tiêu đề. Hãy tải và dùng tệp mẫu.")
        if len(header_cells) > MAX_COLUMNS:
            raise ValidationError("Tệp Excel có quá nhiều cột (tối đa 100).")
        aliases = _aliases(fields)
        exact_labels = {}
        for column in columns:
            exact_labels.setdefault(column["label"].strip().casefold(), set()).add(
                column["key"]
            )
        mapped = {}
        used = set()
        for index, cell in enumerate(header_cells):
            value = cell.value
            if cell.data_type in {"f", "e"}:
                result["errors"].append(
                    f"Tiêu đề cột {index + 1} chứa công thức hoặc lỗi Excel."
                )
                continue
            if value is None or str(value).strip() == "":
                continue
            matches = exact_labels.get(str(value).strip().casefold()) or aliases.get(
                _header(value), set()
            )
            if len(matches) != 1:
                result["errors"].append(
                    f"Không nhận diện được tiêu đề cột {index + 1}: {str(value)[:100]}. Hãy dùng tên cột trong mẫu mới."
                )
                continue
            key = next(iter(matches))
            if key in used:
                result["errors"].append(
                    f"Tiêu đề {str(value)[:100]} xuất hiện nhiều lần."
                )
                continue
            used.add(key)
            mapped[index] = key
        labels = {column["key"]: column["label"] for column in columns}
        missing = REQUIRED_COLUMNS - used
        if missing:
            result["errors"].append(
                "Thiếu cột bắt buộc: "
                + ", ".join(labels[key] for key in sorted(missing))
                + "."
            )
        for number, cells in enumerate(iterator, 2):
            if number > MAX_SCAN_ROWS or len(cells) > MAX_COLUMNS:
                raise ValidationError(
                    "Trang tính quá lớn. Chỉ giữ vùng dữ liệu tối đa 500 mặt hàng và 100 cột."
                )
            if not any(
                cell.value is not None and str(cell.value).strip() for cell in cells
            ):
                continue
            if len(result["rows"]) >= MAX_PRODUCTS:
                raise ValidationError(
                    "Mỗi lần chỉ nhập tối đa 500 mặt hàng. Hãy chia tệp rồi tải lại."
                )
            row = {
                "row_number": number,
                "data": {key: "" for key in labels},
                "errors": [],
            }
            row["data"]["minimum"] = "0"
            row["data"]["qr_auto"] = "True"
            for index, cell in enumerate(cells):
                value = cell.value
                key = mapped.get(index)
                if not key:
                    if value is not None and str(value).strip():
                        _error(
                            row,
                            "__all__",
                            "Cột dữ liệu",
                            f"Cột {index + 1} có dữ liệu nhưng chưa có tiêu đề hợp lệ.",
                        )
                    continue
                if cell.data_type in {"f", "e"}:
                    _error(
                        row,
                        key,
                        labels[key],
                        "Không dùng công thức hoặc ô đang lỗi Excel; hãy dán giá trị thuần.",
                    )
                    continue
                if (
                    key in {"sku", "qr_code", "warehouse_code"}
                    and value is not None
                    and not isinstance(value, str)
                ):
                    _error(
                        row,
                        key,
                        labels[key],
                        "Mã hàng / mã QR / mã kho phải có định dạng văn bản (Text) để giữ nguyên số 0 đầu mã.",
                    )
                text = str(value).strip() if value is not None else ""
                if key == "qr_auto":
                    choice = _header(text)
                    if choice in {"", "co", "yes", "true", "1", "tu dong", "auto"}:
                        text = "True"
                    elif choice in {"khong", "no", "false", "0", "tu nhap", "manual"}:
                        text = "False"
                    else:
                        _error(
                            row, key, labels[key], "Điền Có hoặc Không cho QR tự động."
                        )
                row["data"][key] = text if text or key != "minimum" else "0"
            result["rows"].append(row)
        if not result["rows"]:
            raise ValidationError(
                "Tệp chưa có hàng hóa. Điền dữ liệu từ dòng 2 rồi tải lại."
            )
        if (
            len(json.dumps(result["rows"], ensure_ascii=False).encode())
            > MAX_NORMALIZED_BYTES
        ):
            raise ValidationError(
                "Nội dung hàng hóa quá dài. Hãy chia nhỏ tệp (tối đa 1 MB dữ liệu văn bản mỗi lần)."
            )
        _validate_rows(result["rows"], fields)
        validate_allocations(result["rows"], user)
        validated = True
    except ValidationError as exc:
        result["errors"].extend(exc.messages)
    except Exception:
        result["errors"].append(
            "Không đọc được nội dung Excel. Hãy chép dữ liệu vào tệp mẫu mới rồi tải lại."
        )
    finally:
        if workbook is not None:
            workbook.close()
    for row in result["rows"]:
        row.pop("instance", None)
        row["display"] = dict(row["data"])
    result["allocation_summary"] = allocation_summary(result["rows"])
    result["product_count"] = len({row["data"].get("sku") for row in result["rows"]})
    result["count"] = len(result["rows"])
    result["rows_checked"] = validated
    result["valid_count"] = (
        sum(not row["errors"] for row in result["rows"]) if validated else 0
    )
    result["valid"] = bool(
        result["count"]
        and not result["errors"]
        and result["valid_count"] == result["count"]
    )
    return result


class ProductImportResult(list):
    def __init__(self, products, documents):
        super().__init__(products)
        self.documents = documents


def confirm_product_import(
    user, rows, schema, *, filename="", business_date=None, batch_key=None
):
    """Revalidate server-owned preview data and insert the complete batch atomically."""
    require_function(user, "products_create")
    if not isinstance(rows, list) or not rows or len(rows) > MAX_PRODUCTS:
        raise ValidationError(
            "Danh sách nhập không hợp lệ. Hãy tải lại tệp Excel (1–500 mặt hàng)."
        )
    try:
        payload = json.dumps(rows, ensure_ascii=False).encode()
    except (TypeError, ValueError) as exc:
        raise ValidationError(
            "Dữ liệu xem trước không hợp lệ. Hãy tải lại tệp Excel."
        ) from exc
    # Display/error copies do not count toward the source's one MB text limit.
    if len(payload) > MAX_NORMALIZED_BYTES * 4:
        raise ValidationError("Dữ liệu xem trước quá lớn. Hãy chia nhỏ tệp Excel.")
    try:
        with transaction.atomic():
            from .product_codes import lock_product_codes

            lock_product_codes()
            user = get_user_model().objects.select_for_update().get(pk=user.pk)
            require_function(user, "products_create")
            fields = list(CustomField.objects.select_for_update().order_by("pk"))
            if schema != _schema(fields):
                raise ValidationError(
                    "Mẫu nhập hàng / thông tin bổ sung hàng hóa đã thay đổi sau khi xem trước. Hãy tải mẫu mới và nhập lại."
                )
            keys = {column["key"] for column in _columns(fields)}
            validated = []
            for source in rows:
                if (
                    not isinstance(source, dict)
                    or not isinstance(source.get("data"), dict)
                    or set(source["data"]) - keys
                    or not isinstance(source.get("row_number"), int)
                    or source["row_number"] < 2
                    or any(
                        not isinstance(value, str) for value in source["data"].values()
                    )
                ):
                    raise ValidationError(
                        "Dữ liệu xem trước không hợp lệ. Hãy tải lại tệp Excel."
                    )
                if source.get("errors"):
                    raise ValidationError(
                        "Tệp vẫn còn dòng lỗi. Hãy sửa Excel và tải lại trước khi xác nhận."
                    )
                validated.append(
                    {
                        "row_number": source["row_number"],
                        "data": dict(source["data"]),
                        "errors": [],
                    }
                )
            if (
                len(
                    json.dumps(
                        [row["data"] for row in validated], ensure_ascii=False
                    ).encode()
                )
                > MAX_NORMALIZED_BYTES
            ):
                raise ValidationError(
                    "Nội dung hàng hóa quá dài. Hãy chia nhỏ tệp Excel."
                )
            _validate_rows(validated, fields)
            if any(has_allocation(row) for row in validated):
                require_function(user, "receipts_create")
            validate_allocations(validated, user, lock=True)
            for source, row in zip(rows, validated, strict=True):
                if has_allocation(row) and source.get("warehouse_id") != row.get(
                    "warehouse_id"
                ):
                    _error(
                        row,
                        "warehouse_code",
                        "Mã kho nhập",
                        "Kho đã thay đổi sau bước xem trước. Hãy tải lại file.",
                    )
            failures = [
                f"Dòng {row['row_number']} · {error['label']}: {error['message']}"
                for row in validated
                for error in row["errors"]
            ]
            if failures:
                raise ValidationError(failures)
            products = list(
                {row["instance"].sku: row["instance"] for row in validated}.values()
            )
            Product.objects.bulk_create(products, batch_size=MAX_PRODUCTS)
            documents = create_allocated_receipts(
                user,
                validated,
                products,
                business_date=business_date,
                filename=filename,
                batch_key=batch_key,
            )
            return ProductImportResult(products, documents)
    except IntegrityError as exc:
        raise ValidationError(
            "Có mã hàng vừa được tạo bởi thao tác khác. Chưa nhập hàng nào từ tệp này; hãy tải lại để kiểm tra mã trùng."
        ) from exc
