from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import timedelta
from decimal import Decimal
from threading import Barrier
from unittest import skipUnless
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import close_old_connections, connection, connections
from django.test import TestCase, TransactionTestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import Balance, Document, DocumentEvent, Line, Movement, Product, Warehouse
from .product_import import (
    build_product_template,
    confirm_product_import,
    parse_product_workbook,
)
from .services import post_document
from .stocktakes import freeze_session
from .test_product_import import workbook_upload
from .test_support import grant_functions

HEADERS = [
    "Mã hàng",
    "Tên hàng",
    "Đơn vị tính",
    "Mã kho nhập",
    "Số lượng nhập",
    "Đơn giá nhập",
]


class ProductAllocationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.maker = User.objects.create_user("allocation-maker")
        cls.cataloguer = User.objects.create_user("catalog-only")
        cls.head = User.objects.create_user("allocation-head")
        cls.admin = User.objects.create_superuser(
            "allocation-admin", "owner@example.test", "1"
        )
        grant_functions(cls.maker, {"products_create", "receipts_create"}, level=10)
        grant_functions(cls.cataloguer, {"products_create"})
        grant_functions(cls.head, {"receipts_approve"}, level=30)
        cls.a = Warehouse.objects.create(code="K01", name="Kho A")
        cls.b = Warehouse.objects.create(code="K02", name="Kho B")
        cls.outside = Warehouse.objects.create(code="HIDDEN", name="Ngoài phạm vi")
        for w in [cls.a, cls.b]:
            w.members.add(cls.maker)
            w.managers.add(cls.head)

    def parse(self, rows=None, *, user=None, headers=HEADERS):
        return parse_product_workbook(
            workbook_upload(rows or [["P1", "Hàng 1", "Cái", "K01", 2, 10]], headers),
            user=user or self.maker,
        )

    def confirm(self, preview, **kwargs):
        return confirm_product_import(
            kwargs.pop("user", self.maker),
            preview["rows"],
            preview["schema"],
            filename=preview["filename"],
            business_date=preview["business_date"],
            batch_key="test-batch",
            **kwargs,
        )

    def test_downloaded_template_contains_optional_allocation_and_only_authorized_warehouses(
        self,
    ):
        workbook = load_workbook(build_product_template(user=self.maker))
        labels = [c.value for c in workbook.worksheets[0][1]]
        self.assertTrue(set(HEADERS).issubset(labels))
        warehouse_sheet = workbook["Danh sách kho"]
        self.assertEqual(
            set(r[0] for r in list(warehouse_sheet.values)[1:]), {"K01", "K02"}
        )
        self.assertEqual(
            workbook.worksheets[0]
            .cell(2, labels.index("Mã kho nhập") + 1)
            .number_format,
            "@",
        )
        workbook.close()
        workbook = load_workbook(build_product_template(user=self.cataloguer))
        self.assertEqual(workbook["Danh sách kho"].max_row, 1)
        workbook.close()

    def test_multi_warehouse_preview_then_drafts_then_higher_approval_only(self):
        preview = self.parse(
            [
                ["P1", "Máy A", "Cái", "K01", 2, 10],
                ["P2", "Máy B", "Bộ", "K02", 3, 20],
                ["P3", "Chỉ danh mục", "Cái"],
            ]
        )
        self.assertTrue(preview["valid"], preview)
        self.assertEqual(len(preview["allocation_summary"]), 2)
        self.assertFalse(Product.objects.exists())
        result = self.confirm(preview)
        self.assertEqual((len(result), len(result.documents)), (3, 2))
        self.assertEqual(DocumentEvent.objects.count(), 2)
        self.assertFalse(Balance.objects.exists())
        self.assertFalse(Movement.objects.exists())
        a = Document.objects.get(destination=self.a)
        b = Document.objects.get(destination=self.b)
        self.assertEqual((a.kind, a.status, a.creator_level), ("IN", "DRAFT", 10))
        self.assertEqual(a.lines.get().expected_quantity, Decimal(2))
        with self.assertRaises(PermissionDenied):
            post_document(a.pk, self.maker)
        post_document(a.pk, self.head)
        post_document(b.pk, self.head)
        self.assertEqual(Balance.objects.get(warehouse=self.a).quantity, 2)
        self.assertEqual(Balance.objects.get(warehouse=self.b).value, 60)
        self.assertEqual(Movement.objects.count(), 2)
        self.assertFalse(Line.objects.filter(product__sku="P3").exists())

    def test_same_sku_across_warehouses_creates_one_lifetime_qr(self):
        preview = self.parse(
            [
                ["SHARED", "Cùng mã", "Cái", "K01", 2, 10],
                ["SHARED", "Cùng mã", "Cái", "K02", 3, 15],
            ]
        )
        self.assertTrue(preview["valid"], preview)
        self.assertEqual(preview["product_count"], 1)
        result = self.confirm(preview)
        self.assertEqual(len(result), 1)
        self.assertEqual(Product.objects.count(), 1)
        self.assertEqual(Line.objects.values("product_id").distinct().count(), 1)
        self.assertEqual(Document.objects.count(), 2)
        self.assertTrue(Product.objects.get().qr_code)

    def test_same_manual_qr_across_warehouses_is_one_product_not_a_conflict(self):
        headers = HEADERS + ["QR tự động", "Mã QR"]
        preview = self.parse(
            [
                ["SHARED", "Cùng mã", "Cái", "K01", 2, 10, "Không", "FIXED-QR"],
                ["SHARED", "Cùng mã", "Cái", "K02", 3, 15, "Không", "FIXED-QR"],
            ],
            headers=headers,
        )
        self.assertTrue(preview["valid"], preview)
        self.confirm(preview)
        self.assertEqual(Product.objects.get().qr_code, "FIXED-QR")

    def test_repeated_same_warehouse_or_conflicting_master_information_is_blocked(self):
        for second in [
            ["P1", "Hàng 1", "Cái", "K01", 1, 10],
            ["P1", "Tên khác", "Cái", "K02", 1, 10],
            ["P1", "Hàng 1", "Bộ", "K02", 1, 10],
            ["P1", "Hàng 1", "Cái", "", "", ""],
        ]:
            with self.subTest(second=second):
                preview = self.parse([["P1", "Hàng 1", "Cái", "K01", 1, 10], second])
                self.assertFalse(preview["valid"], preview)
                with self.assertRaises(ValidationError):
                    self.confirm(preview)
        self.assertFalse(Product.objects.exists())
        self.assertFalse(Document.objects.exists())

    def test_incomplete_unknown_or_outside_allocation_rejects_entire_file(self):
        for allocation in [
            ["K01", "", 10],
            ["K01", 1, ""],
            ["", 1, 10],
            ["UNKNOWN", 1, 10],
            ["HIDDEN", 1, 10],
            [1, 1, 10],
        ]:
            with self.subTest(allocation=allocation):
                preview = self.parse(
                    [
                        ["GOOD", "Good", "Cái", "K01", 1, 10],
                        ["BAD", "Bad", "Cái", *allocation],
                    ]
                )
                self.assertFalse(preview["valid"], preview)
                with self.assertRaises(ValidationError):
                    self.confirm(preview)
        self.assertFalse(Product.objects.exists())
        self.assertFalse(Document.objects.exists())

    def test_invalid_quantity_price_formula_and_explicit_free_price(self):
        for quantity, price in [
            (0, 1),
            (-1, 1),
            ("0.0001", 1),
            (1, -1),
            (1, "0.00001"),
            (1, "NaN"),
            ("Infinity", 1),
            (1, "=2+3"),
            ("9999999999999", "99999999999999"),
        ]:
            with self.subTest(quantity=quantity, price=price):
                self.assertFalse(
                    self.parse([["P1", "P", "Cái", "K01", quantity, price]])["valid"]
                )
        preview = self.parse([["P1", "Miễn phí", "Cái", "K01", 1, 0]])
        self.assertTrue(preview["valid"], preview)
        self.confirm(preview)
        self.assertEqual(Line.objects.get().unit_price, 0)

    def test_value_rounding_in_preview_matches_posting(self):
        preview = self.parse([["P1", "P", "Cái", "K01", 1, "0.005"]])
        self.assertEqual(preview["allocation_summary"][0]["amount"], "0.01")
        result = self.confirm(preview)
        post_document(result.documents[0].pk, self.head)
        self.assertEqual(Balance.objects.get().value, Decimal("0.01"))

    def test_catalog_only_user_can_import_master_but_not_allocations(self):
        preview = self.parse(user=self.cataloguer)
        self.assertFalse(preview["valid"])
        preview = self.parse([["P1", "P", "Cái"]], user=self.cataloguer)
        self.assertTrue(preview["valid"], preview)
        result = self.confirm(preview, user=self.cataloguer)
        self.assertEqual(len(result.documents), 0)
        self.assertFalse(Document.objects.exists())

    def test_revoked_receipt_permission_or_warehouse_access_blocks_confirmation(self):
        preview = self.parse()
        grant_functions(self.maker, {"products_create"})
        with self.assertRaises(PermissionDenied):
            self.confirm(preview)
        grant_functions(self.maker, {"products_create", "receipts_create"})
        self.a.members.remove(self.maker)
        with self.assertRaises(ValidationError):
            self.confirm(preview)
        self.assertFalse(Product.objects.exists())
        self.assertFalse(Document.objects.exists())

    def test_same_code_reassigned_to_different_warehouse_is_not_silently_used(self):
        preview = self.parse()
        self.a.code = "RENAMED"
        self.a.save()
        impostor = Warehouse.objects.create(code="K01", name="Kho khác")
        impostor.members.add(self.maker)
        with self.assertRaises(ValidationError):
            self.confirm(preview)
        self.assertFalse(Product.objects.exists())

    def test_day_rollover_and_mid_creation_error_roll_back_all_master_and_drafts(self):
        preview = self.parse(
            [["P1", "P1", "Cái", "K01", 1, 10], ["P2", "P2", "Cái", "K02", 1, 10]]
        )
        preview["business_date"] = str(timezone.localdate() - timedelta(days=1))
        with self.assertRaises(ValidationError):
            self.confirm(preview)
        self.assertFalse(Product.objects.exists())
        preview["business_date"] = str(timezone.localdate())
        from .services import audit_document as real_audit

        audited = []

        def fail_second(doc, actor, action):
            audited.append(doc.pk)
            if len(audited) == 2:
                raise ValidationError("test failure after first receipt")
            return real_audit(doc, actor, action)

        with patch(
            "inventory.product_allocation.audit_document", side_effect=fail_second
        ):
            with self.assertRaises(ValidationError):
                self.confirm(preview)
        self.assertEqual(len(audited), 2)
        self.assertFalse(Product.objects.exists())
        self.assertFalse(Document.objects.exists())
        self.assertFalse(Line.objects.exists())

    def test_51_lines_are_split_into_two_reviewable_receipts(self):
        preview = self.parse(
            [[f"P{i}", f"Hàng {i}", "Cái", "K01", 1, 10] for i in range(51)]
        )
        self.assertTrue(preview["valid"], preview)
        self.assertEqual(preview["allocation_summary"][0]["documents"], 2)
        result = self.confirm(preview)
        self.assertEqual(sorted(d.lines.count() for d in result.documents), [1, 50])
        self.assertEqual(Product.objects.count(), 51)
        self.assertFalse(Balance.objects.exists())

    def test_500_rows_create_ten_drafts_with_bounded_queries(self):
        preview = self.parse(
            [[f"BULK{i}", f"Hàng {i}", "Cái", "K01", 1, 10] for i in range(500)]
        )
        self.assertTrue(preview["valid"], preview["errors"])
        with CaptureQueriesContext(connection) as queries:
            result = self.confirm(preview)
        self.assertLess(len(queries), 130)
        self.assertEqual((len(result), len(result.documents)), (500, 10))
        self.assertTrue(all(doc.lines.count() == 50 for doc in result.documents))
        self.assertFalse(Balance.objects.exists())
        self.assertFalse(Movement.objects.exists())

    def test_freeze_does_not_prevent_draft_but_blocks_approval(self):
        freeze_session(self.a, self.admin)
        result = self.confirm(self.parse())
        with self.assertRaises(ValidationError):
            post_document(result.documents[0].pk, self.head)
        self.assertFalse(Balance.objects.exists())

    def test_http_preview_confirmation_results_scope_and_replay(self):
        self.client.force_login(self.maker)
        response = self.client.post(
            reverse("product_import"),
            {
                "action": "preview",
                "file": workbook_upload(
                    [
                        ["P1", "P1", "Cái", "K01", 1, 10],
                        ["P2", "P2", "Cái", "K02", 1, 20],
                    ],
                    HEADERS,
                ),
            },
        )
        self.assertContains(response, "Phiếu nhập dự kiến theo kho")
        self.assertFalse(Document.objects.exists())
        key = self.client.session["product_import"]["key"]
        self.client.post(
            reverse("product_import"),
            {
                "action": "confirm",
                "import_key": key,
                "warehouse_code": "HIDDEN",
                "receipt_quantity": 999,
            },
        )
        self.assertEqual(Document.objects.count(), 2)
        self.assertEqual(
            set(Document.objects.values_list("destination_id", flat=True)),
            {self.a.pk, self.b.pk},
        )
        self.assertEqual(sum(Line.objects.values_list("quantity", flat=True)), 2)
        response = self.client.get(reverse("products") + "?batch=" + key)
        self.assertContains(response, "2 phiếu nhập đã tạo theo kho")
        self.client.post(
            reverse("product_import"), {"action": "confirm", "import_key": key}
        )
        self.assertEqual(Document.objects.count(), 2)
        self.assertFalse(Balance.objects.exists())
        session = self.client.session
        session["warehouse_scope"] = str(self.a.pk)
        session.save()
        response = self.client.get(
            reverse("documents") + "?kind=IN&import_batch=" + key
        )
        self.assertEqual(response.context["page_obj"].paginator.count, 2)
        self.client.force_login(self.head)
        response = self.client.get(
            reverse("documents") + "?kind=IN&import_batch=" + key
        )
        self.assertEqual(response.context["page_obj"].paginator.count, 0)

    def test_warehouse_pin_missing_in_tampered_preview_is_rejected(self):
        preview = self.parse()
        preview["rows"][0].pop("warehouse_id")
        with self.assertRaises(ValidationError):
            self.confirm(preview)
        self.assertFalse(Product.objects.exists())


@skipUnless(
    connection.vendor == "postgresql", "Requires concurrent PostgreSQL transactions"
)
class ProductAllocationConcurrencyTests(TransactionTestCase):
    def test_concurrent_same_batch_cannot_create_duplicate_products_or_receipts(self):
        user = get_user_model().objects.create_user("allocation-concurrent")
        grant_functions(user, {"products_create", "receipts_create"})
        warehouse = Warehouse.objects.create(code="K01", name="A")
        warehouse.members.add(user)
        preview = parse_product_workbook(
            workbook_upload([["P1", "P1", "Cái", "K01", 1, 10]], HEADERS), user=user
        )
        self.assertTrue(preview["valid"], preview)
        barrier = Barrier(2)

        def create(_):
            close_old_connections()
            try:
                actor = get_user_model().objects.get(pk=user.pk)
                barrier.wait(timeout=10)
                confirm_product_import(
                    actor,
                    deepcopy(preview["rows"]),
                    preview["schema"],
                    business_date=preview["business_date"],
                    batch_key="same-batch",
                )
                return "created"
            except ValidationError:
                return "blocked"
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(create, [0, 1]))
        self.assertEqual(sorted(results), ["blocked", "created"])
        self.assertEqual(Product.objects.count(), 1)
        self.assertEqual(Document.objects.count(), 1)
        self.assertEqual(Line.objects.count(), 1)
        self.assertFalse(Balance.objects.exists())
