from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .issue_import_views import HEADERS, parse_issue
from .models import Balance, Document, Line, Movement, Product, Warehouse
from .services import post_document
from .test_support import grant_functions


def upload(rows, headers=HEADERS):
    book = Workbook()
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return SimpleUploadedFile("xuat.xlsx", buffer.getvalue())


class IssueImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.maker = User.objects.create_user("excel-maker")
        cls.viewer = User.objects.create_user("excel-viewer")
        cls.approver = User.objects.create_user("excel-head")
        grant_functions(cls.maker, {"issues_create", "receipts_create"})
        grant_functions(cls.viewer, {"issues_view"})
        grant_functions(cls.approver, {"issues_approve"}, level=30)
        cls.warehouse = Warehouse.objects.create(code="X", name="Kho Excel")
        cls.other = Warehouse.objects.create(code="Y", name="Kho ngoài quyền")
        cls.warehouse.members.add(cls.maker, cls.viewer)
        cls.warehouse.issue_approvers.add(cls.approver)
        cls.product = Product.objects.create(sku="001A", name="Hàng A", unit="cái")
        Balance.objects.create(
            warehouse=cls.warehouse, product=cls.product, quantity=10, value=250
        )
        cls.url = reverse("issue_import")

    def setUp(self):
        self.client.force_login(self.maker)

    def preview(self, rows=None, **overrides):
        data = {
            "warehouse": self.warehouse.pk,
            "file": upload(rows or [["001A", 3, "Cấp phòng"]]),
            "business_date": str(timezone.localdate()),
            "note": "Cấp vật tư",
        }
        data.update(overrides)
        return self.client.post(self.url, data)

    def confirm(self):
        key = self.client.session["issue_import"]["key"]
        return self.client.post(self.url, {"confirm": 1, "import_key": key})

    def test_navigation_keeps_excel_in_its_business_module(self):
        self.preview()
        from .access import access_context
        from .navigation import navigation

        response = self.client.get(self.url)
        sections = navigation(response.wsgi_request, access_context(self.maker))
        self.assertFalse(
            any(
                item["label"] == "Nhập từ Excel"
                for section in sections
                for item in section["items"]
            )
        )
        self.assertTrue(
            any(
                item["label"] == "Xuất kho" and item["active"]
                for section in sections
                for item in section["items"]
            )
        )
        page = self.client.get(reverse("operation_new", args=["out"]))
        self.assertContains(page, self.url)
        self.assertContains(
            self.client.get(reverse("operation_new", args=["in"])),
            reverse("receipt_import"),
        )

    def test_preview_and_confirmation_do_not_change_stock_then_higher_approval_posts(
        self,
    ):
        response = self.preview()
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "001A")
        self.assertFalse(Document.objects.exists())
        response = self.confirm()
        self.assertEqual(response.status_code, 302)
        doc = Document.objects.get()
        self.assertEqual(
            (doc.kind, doc.status, doc.source_id, doc.destination_id),
            ("OUT", "DRAFT", self.warehouse.pk, None),
        )
        line = doc.lines.get()
        self.assertEqual(line.sku_snapshot, "001A")
        self.assertIsNone(line.unit_price)
        self.assertEqual(Balance.objects.get().quantity, 10)
        self.assertFalse(Movement.objects.exists())
        with self.assertRaises(PermissionDenied):
            post_document(doc.pk, self.maker)
        post_document(doc.pk, self.approver)
        balance = Balance.objects.get()
        self.assertEqual((balance.quantity, balance.value), (Decimal(7), Decimal(175)))
        self.assertEqual(doc.lines.get().book_value, Decimal(75))

    def test_stock_rechecked_after_preview_and_whole_file_rejected(self):
        second = Product.objects.create(sku="B", name="B")
        Balance.objects.create(
            warehouse=self.warehouse, product=second, quantity=5, value=20
        )
        self.preview([["001A", 3], ["B", 5]])
        Balance.objects.filter(product=second).update(quantity=2)
        self.confirm()
        self.assertFalse(Document.objects.exists())
        self.assertFalse(Line.objects.exists())

    def test_preview_does_not_accept_missing_or_insufficient_stock(self):
        self.preview([["001A", 11]])
        self.assertNotIn("issue_import", self.client.session)
        Product.objects.create(sku="LOANONLY", name="Hàng mượn, không có tồn sở hữu")
        self.preview([["LOANONLY", 1]])
        self.assertNotIn("issue_import", self.client.session)

    def test_file_errors_are_atomic(self):
        for rows in [
            [["001A", 1], ["MISSING", 1]],
            [["001A", 1], ["001A", 2]],
            [["001A", "=1+1"]],
            [[1, 2]],
            [["001A", 0]],
            [["001A", -1]],
            [["001A", "NaN"]],
            [["001A", "0.0001"]],
            [["001A", True]],
            [["001A", 1, "a" * 251]],
        ]:
            with self.subTest(rows=rows), self.assertRaises(ValidationError):
                parse_issue(upload(rows))
        self.assertFalse(Document.objects.exists())

    def test_fifty_one_rows_and_wrong_template_rejected(self):
        with self.assertRaises(ValidationError):
            parse_issue(upload([[f"A{i}", 1] for i in range(51)]))
        with self.assertRaises(ValidationError):
            parse_issue(upload([["001A", 1, 20]], ["ma_hang", "so_luong", "don_gia"]))

    def test_replay_or_tampered_post_cannot_duplicate_or_change_preview(self):
        self.preview()
        key = self.client.session["issue_import"]["key"]
        self.client.post(
            self.url,
            {
                "confirm": 1,
                "import_key": key,
                "quantity": 999,
                "warehouse": self.other.pk,
            },
        )
        self.client.post(self.url, {"confirm": 1, "import_key": key})
        self.assertEqual(Document.objects.count(), 1)
        self.assertEqual(Line.objects.get().quantity, 3)
        self.assertEqual(Document.objects.get().source_id, self.warehouse.pk)

    def test_expired_owner_or_wrong_key_and_date_cannot_confirm(self):
        for change in [{"owner": self.viewer.pk}, {"created_at": 0}, {"key": "other"}]:
            self.preview()
            session = self.client.session
            preview = session["issue_import"]
            original = preview["key"]
            preview.update(change)
            session["issue_import"] = preview
            session.save()
            self.client.post(self.url, {"confirm": 1, "import_key": original})
            self.assertFalse(Document.objects.exists())
        self.preview()
        session = self.client.session
        preview = session["issue_import"]
        preview["details"]["business_date"] = str(
            timezone.localdate() - timedelta(days=1)
        )
        session["issue_import"] = preview
        session.save()
        self.confirm()
        self.assertFalse(Document.objects.exists())

    def test_permissions_and_warehouse_rechecked(self):
        self.assertEqual(self.preview(warehouse=self.other.pk).status_code, 200)
        self.assertNotIn("issue_import", self.client.session)
        self.preview()
        self.warehouse.members.remove(self.maker)
        self.assertEqual(self.confirm().status_code, 404)
        self.assertFalse(Document.objects.exists())
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.client.get(reverse("issue_template")).status_code, 403)
        self.assertEqual(self.client.get(reverse("receipt_import")).status_code, 403)

    def test_source_product_change_rejected(self):
        self.preview()
        Product.objects.filter(pk=self.product.pk).update(sku="NEW")
        self.confirm()
        self.assertFalse(Document.objects.exists())

    def test_template_and_receipt_access_follow_creation_permissions(self):
        response = self.client.get(reverse("issue_template"))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        self.assertEqual([c.value for c in book.active[1]], HEADERS)
        book.close()
        self.assertEqual(self.client.get(reverse("receipt_import")).status_code, 200)
